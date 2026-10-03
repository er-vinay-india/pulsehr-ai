"""Replay saved presentation layouts without running generation or updating the DB.

Run with backend/.venv/bin/python. Outputs are diagnostic copies, never certification
of legacy slide claims. Existing content/decision verification remains active.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path
import sqlite3
import sys
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.core import config
from app.services.report_generator import export_spec_to_pptx
from pptx import Presentation


def inspect_geometry(path, slides):
    """Measure object rectangles; these are allocation, not painted-ink coverage."""
    prs = Presentation(path)
    canvas = prs.slide_width.inches * prs.slide_height.inches
    body_area = 11.7 * 4.8  # Current exporter's common content envelope.
    result = []
    for index, (native, spec) in enumerate(zip(prs.slides, slides), 1):
        charts, tables, fonts, outside = [], [], set(), []
        for shape in native.shapes:
            box = dict(x=round(shape.left.inches, 3), y=round(shape.top.inches, 3),
                       width=round(shape.width.inches, 3), height=round(shape.height.inches, 3))
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                outside.append(shape.name)
            if shape.has_chart:
                area = shape.width.inches * shape.height.inches
                charts.append({**box, "canvas_area_pct": round(100 * area / canvas, 1),
                               "common_body_area_pct": round(100 * area / body_area, 1),
                               "series_values": [list(s.values) for s in shape.chart.series]})
            if shape.has_table:
                tables.append({**box, "data_rows": len(shape.table.rows) - 1,
                               "columns": len(shape.table.columns)})
            frames = [shape.text_frame] if shape.has_text_frame else []
            if shape.has_table:
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for frame in frames:
                for paragraph in frame.paragraphs:
                    if paragraph.font.size:
                        fonts.add(round(paragraph.font.size.pt, 1))
                    for run in paragraph.runs:
                        if run.font.size:
                            fonts.add(round(run.font.size.pt, 1))
        result.append({"export_slide": index, "source_slide": spec.get("order"),
                       "title": spec.get("title"), "layout": spec.get("layout"),
                       "explicit_text_font_sizes_pt": sorted(fonts), "charts": charts,
                       "tables": tables, "source_table_data_rows": len((spec.get("table") or {}).get("rows", [])),
                       "missing_native_table": bool(spec.get("table")) and not tables,
                       "outside_canvas_objects": outside})
    return {"canvas_inches": [prs.slide_width.inches, prs.slide_height.inches],
            "measurement_note": "Object allocation only; wrapping, plot interiors and clipping require visual review.",
            "slides": result}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--deck-id", help="Read saved spec from the application DB in read-only mode")
    source.add_argument("--spec", type=Path, help="Replay a previously frozen JSON spec")
    parser.add_argument("--slide", type=int, help="One-based slide to export; default: all")
    parser.add_argument("--layout", choices=["title_cover", "title_hero", "kpi_summary", "chart_narrative",
                        "full_chart_takeaway", "comparison_split", "table_detail", "action_plan"],
                        help="Layout-only experiment on one slide; content fields remain unchanged")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if args.spec:
        raw = args.spec.read_text()
    else:
        with sqlite3.connect(f"{config.DB_PATH.resolve().as_uri()}?mode=ro", uri=True) as conn:
            row = conn.execute("SELECT spec_json FROM presentation_decks WHERE id=?", (args.deck_id,)).fetchone()
        if not row:
            parser.error("Saved deck not found")
        raw = row[0]
    original = json.loads(raw)
    slides = original.get("slides") or []
    if not slides:
        parser.error("Saved deck has no slides")
    if args.slide is not None and not 1 <= args.slide <= len(slides):
        parser.error(f"--slide must be between 1 and {len(slides)}")
    if args.layout and args.slide is None:
        parser.error("--layout requires --slide to keep experiments limited to one slide")
    metadata = original.get("metadata") or {}
    if metadata.get("deck_style") == "decision_brief" and (args.slide or args.layout):
        parser.error("Immutable decision briefs can only be replayed as complete decks")
    if metadata.get("content_contract"):
        from app.services.presentation.content_validation import verify_business_content
        if verify_business_content(original)["status"] != "PASSED":
            parser.error("Source content verification failed; correct the source before testing its layout")
    out = args.output_dir.resolve()
    if out == config.EXPORTS_DIR.resolve() or config.EXPORTS_DIR.resolve() in out.parents:
        parser.error("Choose a diagnostic directory outside the production exports directory")
    out.mkdir(parents=True, exist_ok=True)
    run_id = f"layout_review_{uuid4().hex[:12]}"
    frozen = out / f"{run_id}_source.json"
    frozen.write_text(raw)
    experiment = copy.deepcopy(original)
    experiment["id"] = run_id
    if args.slide:
        selected = args.slide - 1
        experiment["slides"] = [experiment["slides"][selected]]
        # Select existing seals; do not reseal modified wording, values or notes.
        for key in ("content_contract_digests", "content_contract_notes"):
            if key in experiment.get("metadata", {}):
                experiment["metadata"][key] = [experiment["metadata"][key][selected]]
    if args.layout:
        experiment["slides"][0]["layout"] = args.layout
    old_exports = config.EXPORTS_DIR
    try:
        config.EXPORTS_DIR = out
        pptx = export_spec_to_pptx(experiment, diagnostic_layout_only=True)
    finally:
        config.EXPORTS_DIR = old_exports
    from app.services.presentation.slide_layout import resolve_slide
    physical_slides = []
    for slide in experiment['slides']:
        resolved = resolve_slide(slide)
        physical_slides.extend([slide] * (len(resolved['pages']) if resolved else 1))
    report = inspect_geometry(pptx, physical_slides)
    report.update(source_deck_id=original.get("id"), source_sha256=hashlib.sha256(raw.encode()).hexdigest(),
                  diagnostic_only=True, generation_called=False, database_written=False,
                  selected_slide=args.slide, layout_override=args.layout,
                  content_contract=metadata.get("content_contract"), pptx=str(pptx), frozen_source=str(frozen))
    report_path = out / f"{run_id}_geometry.json"
    report_path.write_text(json.dumps(report, indent=2))
    print(json.dumps({"pptx": str(pptx), "geometry": str(report_path), "frozen_source": str(frozen)}, indent=2))


if __name__ == "__main__":
    main()
