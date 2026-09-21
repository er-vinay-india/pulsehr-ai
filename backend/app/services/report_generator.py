"""Executive report and presentation generator service."""

import datetime
import logging
from html import escape
from pathlib import Path
from uuid import uuid4
from pptx import Presentation
from pptx.util import Inches

from ..core import config
from .pptx import (
    C_BG,
    C_CARD,
    C_CARD_BORDER,
    C_BRAND,
    C_ACCENT,
    C_TEXT_LIGHT,
    C_TEXT_MUTED,
    C_SUCCESS,
    C_DANGER,
    ChartExportError,
    hex_to_rgb,
    set_slide_background,
    add_header,
    _save_deck,
    _render_footer,
    _render_title_hero_slide,
    _render_kpi_summary_slide,
    _render_chart_narrative_slide,
    _render_comparison_split_slide,
    _render_table_detail_slide,
    _render_full_chart_takeaway_slide,
    _render_two_charts_slide,
    _render_action_plan_slide,
    _render_generic_slide,
    _table_slide,
)

logger = logging.getLogger(__name__)

__all__ = [
    "ChartExportError",
    "export_spec_to_pptx",
    "generate_pptx_presentation",
    "generate_calculation_presentation",
    "generate_html_executive_report",
    "config",
]


def _new_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _number(value):
    return 'No data' if value is None else f'{value:,.6g}'


def export_spec_to_pptx(deck_spec: dict) -> Path:
    if deck_spec.get("metadata", {}).get("deck_style") == "decision_brief":
        from .presentation.decision_deck import verify_decision_deck
        if verify_decision_deck(deck_spec, deck_spec.get("evidence_ledger", []))["status"] != "PASSED":
            raise ValueError("Decision brief changed after evidence verification. Regenerate the brief before export.")
    prs = _new_deck()
    theme = deck_spec.get("theme", {})
    colors = {
        "bg": hex_to_rgb(theme.get("bg_color", "#171412"), C_BG),
        "card_bg": hex_to_rgb(theme.get("card_bg", "#201b18"), C_CARD),
        "card_border": hex_to_rgb(theme.get("card_border", "#3d362f"), C_CARD_BORDER),
        "primary": hex_to_rgb(theme.get("primary_text", "#fff9f2"), C_TEXT_LIGHT),
        "secondary": hex_to_rgb(theme.get("secondary_text", "#beb2a6"), C_TEXT_MUTED),
        "brand": hex_to_rgb(theme.get("brand_color", "#ff8a62"), C_BRAND),
        "accent": hex_to_rgb(theme.get("accent_color", "#7ee7d9"), C_ACCENT),
        "success": hex_to_rgb(theme.get("success_color", "#8ef0c8"), C_SUCCESS),
        "danger": hex_to_rgb(theme.get("danger_color", "#ff8ca0"), C_DANGER),
    }

    slides = deck_spec.get("slides", [])
    total_slides = len(slides)
    for idx, slide_data in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_slide_background(slide, colors["bg"])
        add_header(
            slide,
            slide_data.get("title", ""),
            slide_data.get("category", "EXECUTIVE REVIEW"),
            brand_color=colors["brand"],
            title_color=colors["primary"]
        )

        layout = slide_data.get("layout", "chart_narrative")
        if layout == "title_hero":
            _render_title_hero_slide(slide, slide_data, colors)
        elif layout == "kpi_summary":
            _render_kpi_summary_slide(slide, slide_data, colors)
        elif layout == "chart_narrative":
            _render_chart_narrative_slide(slide, slide_data, colors, theme_palette=theme.get("chart_palette"))
        elif layout == "full_chart_takeaway":
            _render_full_chart_takeaway_slide(slide, slide_data, colors, theme_palette=theme.get("chart_palette"))
        elif layout == "two_charts":
            _render_two_charts_slide(slide, slide_data, colors, theme_palette=theme.get("chart_palette"))
        elif layout in ("action_plan", "initiative_detail"):
            _render_action_plan_slide(slide, slide_data, colors)
        elif layout in ("comparison_split", "methodology_panel"):
            _render_comparison_split_slide(slide, slide_data, colors)
        elif layout == "table_detail":
            _render_table_detail_slide(slide, slide_data, colors)
        else:
            _render_generic_slide(slide, slide_data, colors)

        _render_footer(slide, slide_data, colors, slide_num=idx + 1, total_slides=total_slides)

        if slide_data.get("speaker_notes"):
            try:
                notes_slide = slide.notes_slide
                tf = notes_slide.notes_text_frame
                tf.text = str(slide_data["speaker_notes"])
            except Exception:
                pass

    deck_id = deck_spec.get("id", uuid4().hex)
    out_path = config.EXPORTS_DIR / f"presentation_{deck_id}.pptx"
    prs.save(str(out_path))
    return out_path


def generate_pptx_presentation() -> Path:
    from .sheet_catalog import overview
    data = overview()
    if not data['sheets']:
        raise ValueError('Upload a sheet before generating a report.')
    prs = _new_deck()
    stamp = f"Generated {datetime.datetime.now():%Y-%m-%d %H:%M}. Source rows are not unique employees."
    _table_slide(prs, 'Uploaded data overview', ['Metric', 'Value'], [
        ['Datasets', data['stats']['datasets']], ['Sheets', data['stats']['sheets']],
        ['Source rows', data['stats']['rows']], ['Exact key relationships', data['stats']['linked_relationships']],
    ], stamp)
    for sheet in data['sheets']:
        metrics = [[p['column'] + (f" ({p['unit']})" if p.get('unit') else ''), p['nonempty'], p['missing'], _number(p.get('numeric', {}).get('mean'))] for p in sheet['profiles']]
        for start in range(0, max(1, len(metrics)), 8):
            _table_slide(prs, sheet['name'][:65], ['Column', 'Present', 'Missing', 'Mean'], metrics[start:start+8],
                         f"Source: {sheet['original_name']} / {sheet['name']}. {sheet['row_count']} rows. Means use numeric, nonmissing cells only.")
    return _save_deck(prs)


def generate_calculation_presentation(result: dict) -> Path:
    prs = _new_deck()
    rows = [[item['group'], _number(item['value']), item['used_rows'], item['missing_rows']] for item in result['results']]
    stamp = f"Source: {result['source']}. Selected {result['matched_rows']} of {result['source_rows']} rows. {result['note']}"
    for start in range(0, max(1, len(rows)), 8):
        _table_slide(prs, f"{result['operation'].title()}: {result['column'] or 'row count'}", ['Group', 'Value', 'Rows used', 'Missing'], rows[start:start+8], stamp)
    _table_slide(prs, 'Calculation scope', ['Setting', 'Value'], [
        ['Source', result['source']], ['Operation', result['operation']],
        ['Column', result['column'] or 'Rows'], ['Group by', result['group_by'] or 'None'],
        ['Filter column', result['filter_column'] or 'None'], ['Filter value', result['filter_value'] or 'None'],
    ], 'Calculated locally from complete source data.')
    return _save_deck(prs)


def generate_html_executive_report() -> str:
    from .sheet_catalog import overview
    data = overview()
    sections = []
    for sheet in data['sheets']:
        rows = ''.join(f"<tr><td>{escape(p['column'])}</td><td>{p['nonempty']}</td><td>{p['missing']}</td><td>{_number(p.get('numeric', {}).get('mean'))}</td></tr>" for p in sheet['profiles'])
        sections.append(f"<h2>{escape(sheet['original_name'])} / {escape(sheet['name'])}</h2><p>{sheet['row_count']} source rows</p><table><tr><th>Column</th><th>Present</th><th>Missing</th><th>Mean</th></tr>{rows}</table>")
    return '<!doctype html><html><head><title>PulseHR AI Report</title><style>body{font-family:Arial;padding:32px}table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccc;padding:8px;text-align:left}</style></head><body><h1>PulseHR AI · Uploaded data report</h1><p>' + escape(data['note']) + '</p>' + (''.join(sections) or '<p>No sheets uploaded.</p>') + '</body></html>'
