"""Assert colours in the actual native PPTX package, not just preset declarations."""
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pytest
from PIL import Image
from app.core import config
from app.services.report_generator import export_spec_to_pptx
from app.services.presentation.visual.design_tokens import SLIDE_THEME_PRESETS, slide_theme_preset
from app.services.presentation.photo_background import _contrast, photo_scrim_opacity, add_photo_background
from pptx import Presentation

NS = {'a':'http://schemas.openxmlformats.org/drawingml/2006/main', 'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}

@pytest.mark.parametrize('theme_id', list(SLIDE_THEME_PRESETS))
def test_native_export_contrast(theme_id, tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'EXPORTS_DIR', tmp_path)
    chart = {'categories':['North','South'], 'series':[{'name':'Attendance','values':[10,20]}]}
    slides = [{'title':'Contrast audit','subtitle':'Audited sample','narrative':'[Evidence] **Verified** result','layout':'title_cover','bullets':['Recommendation']},
              {'title':'Chart labels','layout':'chart_narrative','chart':chart},
              {'title':'Chart segments','layout':'chart_narrative','chart':{**chart,'type':'donut'}},
              {'title':'Table boundaries','layout':'table_detail','table':{'headers':['Metric','Value'],'rows':[['Coverage','97%']]}}]
    # metadata-only deck and stale saved foreground values must both use the repaired preset.
    spec = {'metadata':{'theme_id':theme_id}, 'theme':{'secondary_text':'#888888'}, 'slides':slides}
    path = export_spec_to_pptx(spec)
    theme = slide_theme_preset(theme_id)
    with ZipFile(path) as package:
        text_count = 0
        for name in package.namelist():
            if not (name.startswith('ppt/slides/slide') or name.startswith('ppt/charts/chart')) or not name.endswith('.xml'):
                continue
            root=ET.fromstring(package.read(name))
            for properties in root.findall('.//a:rPr',NS)+root.findall('.//a:defRPr',NS):
                color=properties.find('a:solidFill/a:srgbClr',NS)
                if color is None: continue
                foreground='#'+color.attrib['val'];text_count+=1
                for bg in (theme['bg_color'],theme['card_bg']): assert _contrast(foreground,bg)>=7,(theme_id,name,foreground,bg)
            if name.startswith('ppt/charts/chart'):
                fills=root.findall('c:spPr/a:solidFill/a:srgbClr',NS)+root.findall('c:chart/c:plotArea/c:spPr/a:solidFill/a:srgbClr',NS)
                assert len(fills)==2
                assert all('#'+c.attrib['val'].lower()==theme['card_bg'].lower() for c in fills)
                if root.find('c:chart/c:legend', NS) is not None:
                    assert root.find('c:chart/c:legend/c:txPr', NS) is not None
        assert text_count > 15
    prs=Presentation(path)
    assert len(prs.slides)==len(slides)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart and shape.chart.chart_type not in (5,-4120):
                # Pie/donut do not have axes. Cartesian axes and legends are explicitly coloured.
                try:
                    axis=shape.chart.category_axis
                except ValueError:
                    continue
                assert str(axis.tick_labels.font.color.rgb).lower()==theme['secondary_text'].lstrip('#').lower()
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for edge in ('lnL','lnR','lnT','lnB'):
                            color=cell._tc.find(f'a:tcPr/a:{edge}/a:solidFill/a:srgbClr',NS)
                            assert color is not None
                            assert _contrast('#'+color.attrib['val'],theme['card_bg'])>=3

@pytest.mark.parametrize('theme_id', list(SLIDE_THEME_PRESETS))
def test_photo_extremes_have_readable_text(theme_id, monkeypatch):
    theme=slide_theme_preset(theme_id)
    foregrounds=[theme[k] for k in ('primary_text','secondary_text','muted_text','brand_color','accent_color','success_color','danger_color','warning_color')]
    for pixel in (0,255):
        source=Image.new('RGB',(16,9),(pixel,)*3);data=BytesIO();source.save(data,format='PNG')
        monkeypatch.setattr('app.services.presentation.photo_background.photo_bytes',lambda url:data.getvalue())
        prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
        add_photo_background(slide,{'background_image':'https://example.test/photo','scrim_opacity':0},prs.slide_width,prs.slide_height,theme)
        image=Image.open(BytesIO(slide.shapes[0].image.blob));bg='#'+''.join(f'{v:02x}' for v in image.getpixel((100,100)))
        assert all(_contrast(fg,bg)>=7 for fg in foregrounds)
        assert photo_scrim_opacity({'scrim_opacity':100},theme)==1

def test_saved_deck_download_rebuilds_stale_file(tmp_path, monkeypatch):
    import json
    from contextlib import contextmanager
    from app.routers import presentations
    monkeypatch.setattr(config, 'EXPORTS_DIR', tmp_path)
    stale = tmp_path / 'presentation_saved.pptx'
    stale.write_bytes(b'old cached palette')
    spec = {'id':'saved','metadata':{'title':'Readable deck','theme_id':'clean_light'},'slides':[]}
    class Connection:
        def execute(self, *args): return self
        def fetchone(self): return {'spec_json':json.dumps(spec)}
    @contextmanager
    def connection(): yield Connection()
    monkeypatch.setattr(presentations,'get_connection',connection)
    calls=[]
    def export(deck):
        calls.append(deck);stale.write_bytes(b'rebuilt with current palette');return stale
    monkeypatch.setattr(presentations,'export_spec_to_pptx',export)
    response=presentations.download_deck_pptx('saved')
    assert calls == [spec]
    assert response.path == str(stale)
    assert stale.read_bytes() == b'rebuilt with current palette'
