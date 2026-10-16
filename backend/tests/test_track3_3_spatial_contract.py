"""Product Track 3.3 — Spatial Contract, Pagination & Slide Budgeting Acceptance Suite.

Verifies all Track 3.3 Final Spatial Acceptance Closure requirements:
1. Adaptive Geometry-Driven Ranking Pagination (removes fixed 10-items-per-page policy):
   - Same item count + short labels -> fewer pages than same item count + long labels
   - 5 extremely long category labels -> 2 pages
   - 30 long-label ranking items -> multi-page continuation preserving global rank order (#1..#30)
2. Visible Narrative Content Conservation:
   - 200-word, 400-word, and 800-word narratives conserved 100% across visible slide narratives
   - 0 missing words, 0 duplicate words, 0 overflow narrative dumped into speaker_notes
3. Four Critical Fixtures (A, B, C, D) & Three-Consumer One-Pass Page-Boundary Parity:
   - Planned physical slides == Studio resolved pages == PPTX slides == PDF pages
   - 0 additional pagination decisions by Studio, PPTX, or PDF
4. PhysicalSlideContract Schema & Lineage / Hash Separation:
   - physical_slide_id, topic_id, fact_id, evidence_ids, semantic_hash, physical_plan_hash,
     continuation_index, continuation_count, layout_family, content_slice, geometry_certification
5. Canonical 24/24 Semantic Immutability After Spatial Planning:
   - Workforce (8), Retail (8), Environmental (8) -> 24/24 exact match on all 11 protected attributes + semantic_hash
6. Complete Stress Coverage Matrix:
   - Table row counts: 8, 15, 25, 50, 100
   - Table column counts: 4, 8, 12, 20
   - Chart category counts: 5, 10, 15, 25, 50 (short and very long labels)
   - Image shapes: landscape, portrait, square, very_wide, very_tall
   - Narrative lengths: 50, 100, 200, 400, 800 words
7. Zero SlideLayoutError Escape & Typed SpatialPlanningError Escalation
8. Deterministic Spatial Performance & Zero Model Calls
"""

from __future__ import annotations

import re
import statistics
import time
import zipfile
from pathlib import Path
from unittest.mock import patch
import pytest
from pptx import Presentation

from app.db.database import get_connection


def _count_pdf_pages(pdf_path: Path) -> int:
    data = pdf_path.read_bytes()
    return len(re.findall(rb"/Type\s*/Page\b", data))
from app.services.presentation.topic_contract import (
    PresentationTopic,
    compute_canonical_semantic_hash,
)
from app.services.presentation.spatial_planner import (
    PhysicalSlideContract,
    SpatialCanvasBudgeter,
    SpatialPlanningError,
    apply_spatial_slide_budgeting,
)
from app.services.presentation.slide_layout import resolve_slide, SlideLayoutError
from app.services.presentation.resolved_pdf import export_resolved_pdf
from app.services.presentation.deck_generator import generate_presentation_deck_spec
from app.services.presentation.scope_detector import capture_dataset_context, collect_workspace_evidence
from app.services.report_generator import export_spec_to_pptx

pytestmark = pytest.mark.presentation

REQUIRED_CONTRACT_FIELDS = {
    "physical_slide_id",
    "topic_id",
    "fact_id",
    "evidence_ids",
    "semantic_hash",
    "physical_plan_hash",
    "continuation_index",
    "continuation_count",
    "layout_family",
    "content_slice",
    "geometry_certification",
}


def _assert_physical_slide_contract(ps: dict, expected_count: int | None = None):
    missing = REQUIRED_CONTRACT_FIELDS - set(ps.keys())
    assert not missing, f"Missing PhysicalSlideContract fields: {missing}"
    validated = PhysicalSlideContract.model_validate(ps)
    assert validated.geometry_certification.get("certified") is True
    assert validated.geometry_certification.get("zero_overflow") is True
    assert validated.geometry_certification.get("zero_collision") is True
    assert validated.geometry_certification.get("resolved_page_count") == 1
    if expected_count is not None:
        assert validated.continuation_count == expected_count


def _build_word_narrative(word_count: int) -> str:
    base_sentences = [
        "Empirical operational telemetry confirms significant throughput variance across regional distribution facilities.",
        "Primary fulfillment hubs maintained baseline cycle efficiency within governed service tolerances during peak load.",
        "Secondary logistics nodes experienced localized queue latency driven by seasonal shift transitions and dock utilization.",
        "Cross-functional diagnostic audits indicate scheduled preventive maintenance reduced unscheduled downtime across plants.",
        "Capital allocation toward automated sensor calibration stabilizes long-term quality compliance and unit economics.",
    ]
    words: list[str] = []
    idx = 0
    while len(words) < word_count:
        sentence_words = f"Sentence-{idx+1}: {base_sentences[idx % len(base_sentences)]}".split()
        need = word_count - len(words)
        if need < len(sentence_words):
            partial = sentence_words[:need]
            if not partial[-1].endswith("."):
                partial[-1] = partial[-1].rstrip(",;:") + "."
            words.extend(partial)
        else:
            words.extend(sentence_words)
        idx += 1
    assert len(words) == word_count
    return " ".join(words)


@pytest.fixture(scope="function", autouse=True)
def seed_benchmark_datasets():
    """Populates the isolated test SQLite database with the 3 benchmark datasets from production DB."""
    real_db = Path(__file__).resolve().parents[1] / "data" / "db" / "pulsehr.sqlite3"
    with get_connection() as conn:
        has_99770 = conn.execute("SELECT 1 FROM dataset_uploads WHERE id = 99770").fetchone()
        if not has_99770 and real_db.exists():
            conn.execute(f'ATTACH DATABASE "{real_db}" AS real_source')
            conn.execute(
                "INSERT OR REPLACE INTO dataset_uploads (id, filename, original_name, file_type, display_name) "
                "SELECT id, filename, original_name, file_type, display_name FROM real_source.dataset_uploads WHERE id IN (99770, 99771, 99772)"
            )
            conn.execute(
                "INSERT OR REPLACE INTO sheets (id, dataset_id, name, display_name, columns_json, profile_json, row_count) "
                "SELECT id, dataset_id, name, display_name, columns_json, profile_json, row_count FROM real_source.sheets WHERE dataset_id IN (99770, 99771, 99772)"
            )
            conn.execute(
                "INSERT OR REPLACE INTO sheet_rows (id, sheet_id, row_index, data_json) "
                "SELECT id, sheet_id, row_index, data_json FROM real_source.sheet_rows WHERE sheet_id IN (SELECT id FROM real_source.sheets WHERE dataset_id IN (99770, 99771, 99772))"
            )
            try:
                conn.execute(
                    "INSERT OR REPLACE INTO sheet_relationships SELECT * FROM real_source.sheet_relationships "
                    "WHERE left_sheet IN (SELECT id FROM real_source.sheets WHERE dataset_id IN (99770, 99771, 99772))"
                )
            except Exception:
                pass
            conn.commit()
            conn.execute("DETACH DATABASE real_source")


# ==============================================================================
# 1. ADAPTIVE RANKING PAGINATION (GEOMETRY-DRIVEN, NOT FIXED 10)
# ==============================================================================

def test_adaptive_ranking_pagination_geometry_driven():
    """Proves ranking pagination is driven by measured label & plot geometry, not fixed chunks of 10.

    Verifies:
    - 12 short category labels fit on 1 physical page (fewer pages than 12 long labels).
    - 12 long category labels (2-3 wrapped lines) require 2-3 physical pages.
    - 5 extremely long category labels (4 wrapped lines) paginate across 2 physical pages.
    """
    values_12 = [float(100 - i * 3) for i in range(12)]
    short_cats_12 = [f"Node {i+1:02d}" for i in range(12)]
    long_cats_12 = [
        f"Regional Manufacturing & Automated Distribution Facility Unit #{i+1:02d} (North-West Sector)"
        for i in range(12)
    ]

    slide_short_12 = {
        "id": "slide_rank_short_12",
        "topic_id": "TOPIC-RANK-SHORT-12",
        "fact_id": "FACT-RANK-12",
        "evidence_ids": ["EVID-R12"],
        "layout": "chart_narrative",
        "title": "Regional Throughput Ranking",
        "subtitle": "12 Operational Nodes",
        "narrative": "Short-label ranking fits compactly on one physical page.",
        "chart": {
            "type": "ranked_bar",
            "categories": short_cats_12,
            "series": [{"name": "Throughput", "values": values_12}],
            "reference_line": 80.0,
        },
        "semantic_hash": "hash_short_12",
    }

    slide_long_12 = {
        **slide_short_12,
        "id": "slide_rank_long_12",
        "chart": {
            "type": "ranked_bar",
            "categories": long_cats_12,
            "series": [{"name": "Throughput", "values": values_12}],
            "reference_line": 80.0,
        },
    }

    planned_short_12 = SpatialCanvasBudgeter.plan_slide_continuations(slide_short_12)
    planned_long_12 = SpatialCanvasBudgeter.plan_slide_continuations(slide_long_12)

    assert len(planned_short_12) == 1, f"12 short labels should fit on 1 page, got {len(planned_short_12)}"
    assert len(planned_long_12) >= 2, f"12 long multi-line labels must paginate into >=2 pages, got {len(planned_long_12)}"
    assert len(planned_short_12) < len(planned_long_12)

    # 5 extremely long category labels (4 wrapped lines each at 12pt in gutter) -> 2 pages
    ultra_long_cats_5 = [
        f"Global Consolidated Intercontinental Logistics, Cold-Chain Biopharma Warehousing & Automated Customs Hub #{i+1:02d}"
        for i in range(5)
    ]
    slide_ultra_5 = {
        **slide_short_12,
        "id": "slide_rank_ultra_5",
        "chart": {
            "type": "ranked_bar",
            "categories": ultra_long_cats_5,
            "series": [{"name": "Throughput", "values": [95.0, 91.0, 87.0, 82.0, 78.0]}],
            "reference_line": 85.0,
        },
    }
    planned_ultra_5 = SpatialCanvasBudgeter.plan_slide_continuations(slide_ultra_5)
    assert len(planned_ultra_5) == 2, f"5 ultra-long category labels must paginate into 2 pages, got {len(planned_ultra_5)}"
    reconstructed_5 = [c for ps in planned_ultra_5 for c in ps["chart"]["categories"]]
    assert reconstructed_5 == ultra_long_cats_5


# ==============================================================================
# 2. VISIBLE NARRATIVE CONTENT CONSERVATION (200, 400, 800 WORDS)
# ==============================================================================

@pytest.mark.parametrize("word_count", [200, 400, 800])
@pytest.mark.parametrize("with_chart", [False, True])
def test_visible_narrative_conservation_200_400_800_words(word_count: int, with_chart: bool):
    """Proves visible narrative is 100% conserved across continuation slides and never dumped into speaker_notes."""
    narrative_text = _build_word_narrative(word_count)
    original_notes = "Presenter Lineage Note: TOPIC-NARR-CONSERVATION"

    raw_slide = {
        "id": f"slide_narr_{word_count}_{with_chart}",
        "topic_id": f"TOPIC-NARR-{word_count}",
        "fact_id": "FACT-NARR-01",
        "evidence_ids": ["EVID-NARR-01"],
        "layout": "chart_narrative" if with_chart else "comparison_split",
        "title": f"Executive Strategic Narrative Conservation Proof ({word_count} Words)",
        "subtitle": "Audited Multi-Page Visible Narrative Flow",
        "narrative": narrative_text,
        "speaker_notes": original_notes,
        "semantic_hash": f"sem_hash_{word_count}",
    }
    if with_chart:
        raw_slide["chart"] = {
            "type": "ranked_bar",
            "categories": ["Hub Alpha", "Hub Beta", "Hub Gamma", "Hub Delta"],
            "series": [{"name": "Efficiency", "values": [94.2, 89.5, 84.1, 79.8]}],
            "reference_line": 85.0,
        }

    planned = SpatialCanvasBudgeter.plan_slide_continuations(raw_slide)
    assert len(planned) >= 2, f"{word_count}-word narrative must paginate across multiple visible slides"

    # 1. Verify speaker_notes was NOT used as an overflow dump
    for ps in planned:
        assert ps.get("speaker_notes") == original_notes, "Overflow narrative must NEVER be dumped into speaker_notes!"
        _assert_physical_slide_contract(ps, expected_count=len(planned))

    # 2. Verify exact word/token conservation across visible slide narratives
    source_words = narrative_text.split()
    visible_slide_words: list[str] = []
    for ps in planned:
        vis_narr = (ps.get("narrative") or "").strip()
        assert vis_narr, "Every narrative continuation page must carry visible narrative text"
        # Also verify resolve_slide places this exact text into a visible 'narrative' text block on the canvas
        resolved = resolve_slide(ps)
        assert len(resolved["pages"]) == 1
        narrative_blocks = [b for b in resolved["blocks"] if b.get("field") == "narrative"]
        assert len(narrative_blocks) == 1, "Visible narrative must be rendered as a canvas text block, not relegated to notes"
        assert narrative_blocks[0]["text"] == vis_narr
        visible_slide_words.extend(vis_narr.split())

    assert visible_slide_words == source_words, (
        f"Visible narrative token conservation failed for {word_count} words: "
        f"{len(visible_slide_words)} != {len(source_words)}"
    )

    # 3. Verify sentence conservation
    source_sentences = [s.strip() for s in narrative_text.split(".") if s.strip()]
    reconstructed_sentences = [s.strip() for s in " ".join(visible_slide_words).split(".") if s.strip()]
    assert reconstructed_sentences == source_sentences


# ==============================================================================
# 3. FOUR CRITICAL FIXTURES & 3-CONSUMER PAGE-BOUNDARY PARITY
# ==============================================================================

def test_fixture_1_image_and_table_mixed_media_continuation(tmp_path: Path):
    """Fixture B: Image + 25-row Table + Caption + Footer on the same logical topic."""
    headers = ["Facility ID", "Turbine Temperature", "Vibration Score", "Compliance Status"]
    rows = [
        [f"FAC-TURB-{i+1:03d}", f"{65.2 + i * 1.1:.1f}°C", f"{0.12 + i * 0.02:.2f}g", "Optimal" if i % 3 != 0 else "Review"]
        for i in range(25)
    ]

    exec_topic = {
        "topic_id": "TOPIC-MIXED-MEDIA-01",
        "title": "Comprehensive High-Precision Turbine Thermal Inspection & Audit",
        "subtitle": "Predictive Diagnostic Asset Health Review",
        "key_metric": "25 Turbines Audited",
        "analytical_intent": "DIAGNOSTIC",
        "recommended_visual": "table",
        "primary_takeaway": "Thermal anomaly detected in Facility FAC-TURB-018 requiring immediate maintenance window.",
        "inspect_payload": {
            "fact_id": "FACT-TURBINE-ANOMALY-18",
            "evidence_ids": ["EVID-IOT-001", "EVID-IOT-002"],
            "leader_entity": "FAC-TURB-018",
            "leader_value": 84.9,
        },
        "visual_spec": {
            "chart_type": "table",
            "headers": headers,
            "rows": rows,
        },
    }

    pt = PresentationTopic.from_executive_topic(exec_topic, domain="engineering")
    expected_semantic_hash = pt.canonical_semantic_hash()

    raw_slide = pt.to_slide_contract(slide_order=1, total_slides=1)
    raw_slide["image_url"] = "https://assets.highview.ai/facilities/landscape_plant_01.png"
    raw_slide["image_shape"] = "landscape"
    raw_slide["caption"] = "High-resolution thermal infrared scan of turbine assembly."

    planned_slides = SpatialCanvasBudgeter.plan_slide_continuations(raw_slide, theme_id="executive_dark")
    assert len(planned_slides) >= 3, "Image + 25-row table must paginate into 1 image overview + multiple table pages"

    # Part 1: Image Overview Slide
    p1 = planned_slides[0]
    assert p1["continuation_index"] == 1
    assert p1["continuation_count"] == len(planned_slides)
    assert f"Part 1 of {len(planned_slides)}" in p1["title"]
    assert p1["layout"] == "comparison_split"
    assert p1["image_url"] == "https://assets.highview.ai/facilities/landscape_plant_01.png"
    assert p1["table"] is None

    # Parts 2..N: Table Continuation Slides with repeated headers & 100% row conservation
    table_parts = planned_slides[1:]
    collected_rows = []
    for tp in table_parts:
        assert tp["layout"] == "table_detail"
        assert tp["image_url"] is None
        assert tp["table"]["headers"] == headers
        collected_rows.extend(tp["table"]["rows"])

    assert collected_rows == rows, "All 25 rows must be conserved in exact order"

    # Lineage & Hash invariants
    for ps in planned_slides:
        _assert_physical_slide_contract(ps, expected_count=len(planned_slides))
        assert ps["topic_id"] == "TOPIC-MIXED-MEDIA-01"
        assert ps["fact_id"] == "FACT-TURBINE-ANOMALY-18"
        assert ps["evidence_ids"] == ["EVID-IOT-001", "EVID-IOT-002"]
        assert ps["semantic_hash"] == expected_semantic_hash

    plan_hashes = [ps["physical_plan_hash"] for ps in planned_slides]
    assert len(plan_hashes) == len(set(plan_hashes))


def test_fixture_2_long_title_and_narrative_and_chart_pre_budgeting(tmp_path: Path):
    """Fixture A: Long Title + 200+ Word Narrative + Chart."""
    long_title = (
        "Comprehensive Long-Term Strategic Operational Review Analyzing Cross-Functional "
        "Interdependencies Across Workforce Demographics, Longitudinal Attendance Trends, "
        "Regional Performance Variances, and Multi-Quarter Productivity Dynamics in High-Growth Environments"
    )
    long_narrative = _build_word_narrative(220)

    raw_slide = {
        "id": "slide_test_long",
        "topic_id": "TOPIC-LONG-TEST-02",
        "fact_id": "FACT-OPERATIONAL-LATENCY",
        "evidence_ids": ["EVID-OPS-001"],
        "layout": "chart_narrative",
        "title": long_title,
        "subtitle": "Cross-Functional Telemetry Analysis",
        "narrative": long_narrative,
        "primary_takeaway": "Maintenance windows reduced unscheduled downtime by 34%.",
        "chart": {
            "type": "bar",
            "categories": ["Region North", "Region South", "Region East", "Region West"],
            "series": [{"name": "Throughput Index", "values": [88.4, 76.2, 91.5, 68.9]}],
            "reference_line": 80.0,
        },
        "semantic_hash": compute_canonical_semantic_hash({
            "topic_id": "TOPIC-LONG-TEST-02",
            "fact_id": "FACT-OPERATIONAL-LATENCY",
            "evidence_ids": ["EVID-OPS-001"],
            "key_metric": "34%",
            "analytical_intent": "TREND",
            "visual_spec": {
                "chart_type": "bar",
                "categories": ["Region North", "Region South", "Region East", "Region West"],
                "values": [88.4, 76.2, 91.5, 68.9],
            },
        }),
    }

    with pytest.raises(SlideLayoutError):
        resolve_slide(raw_slide)

    budgeted_slides = SpatialCanvasBudgeter.plan_slide_continuations(raw_slide, theme_id="executive_dark")
    assert len(budgeted_slides) >= 2

    reconstructed_words = [w for ps in budgeted_slides for w in (ps.get("narrative") or "").split()]
    assert reconstructed_words == long_narrative.split()


def test_fixture_3_wide_table_identity_safety():
    """Fixture C: Wide table with 8 columns repeating identifying column ('Department') on every part."""
    headers = ["Department", "Headcount", "Active FTE", "Total Attendance", "Leave Hours", "Overtime Hours", "Budget $", "Variance %"]
    rows = [
        ["Engineering", "120", "118", "94.2%", "34h", "12h", "$1,450,000", "+3.2%"],
        ["Product", "45", "44", "91.8%", "18h", "4h", "$620,000", "-1.1%"],
        ["Design", "25", "25", "96.0%", "10h", "2h", "$340,000", "+0.4%"],
        ["Marketing", "35", "33", "88.5%", "22h", "6h", "$480,000", "+2.8%"],
        ["Operations", "80", "78", "92.1%", "28h", "8h", "$950,000", "-0.5%"],
    ]

    raw_slide = {
        "id": "slide_wide_table",
        "topic_id": "TOPIC-WIDE-TABLE-03",
        "fact_id": "FACT-DEPT-METRICS",
        "evidence_ids": ["EVID-HR-001"],
        "layout": "table_detail",
        "title": "Enterprise Departmental Workforce & Financial Audit",
        "subtitle": "Q3 Comprehensive Cross-Functional Ledger",
        "narrative": "Detailed audit of headcount, attendance rates, and operational variance by business unit.",
        "table": {"headers": headers, "rows": rows},
        "semantic_hash": compute_canonical_semantic_hash({
            "topic_id": "TOPIC-WIDE-TABLE-03",
            "fact_id": "FACT-DEPT-METRICS",
            "evidence_ids": ["EVID-HR-001"],
            "key_metric": "$3,840,000",
            "analytical_intent": "AUDIT",
            "visual_spec": {"chart_type": "table", "headers": headers, "rows": rows},
        }),
    }

    budgeted_slides = SpatialCanvasBudgeter.plan_slide_continuations(raw_slide, theme_id="executive_dark")
    assert len(budgeted_slides) == 3, "8-column table (1 ID + 7 value cols) splits into 3 parts of (1 ID + up to 3 value cols)"

    collected_val_headers = []
    for ps in budgeted_slides:
        _assert_physical_slide_contract(ps, expected_count=3)
        p_table = ps["table"]
        assert p_table["headers"][0] == "Department"
        assert len(p_table["rows"]) == len(rows)
        collected_val_headers.extend(p_table["headers"][1:])

    assert collected_val_headers == headers[1:]


def test_fixture_4_large_ranking_chart_global_order_preservation():
    """Fixture D: Large Ranking Chart with 30 Entities and Long Labels."""
    categories = [f"Strategic Business Unit {i+1:02d} — High-Volume Operations Node" for i in range(30)]
    values = [float(100.0 - i * 2.5) for i in range(30)]

    raw_slide = {
        "id": "slide_ranking_30",
        "topic_id": "TOPIC-RANKING-30",
        "fact_id": "FACT-GLOBAL-RANK-30",
        "evidence_ids": ["EVID-RANK-001"],
        "layout": "chart_narrative",
        "title": "Comprehensive Global Business Unit Productivity Ranking",
        "subtitle": "Performance Index Across All 30 Enterprise Nodes",
        "narrative": "Global ranking of all 30 enterprise operational nodes against corporate productivity benchmark.",
        "chart": {
            "type": "ranked_bar",
            "categories": categories,
            "series": [{"name": "Productivity Index", "values": values}],
            "reference_line": 75.0,
        },
        "semantic_hash": compute_canonical_semantic_hash({
            "topic_id": "TOPIC-RANKING-30",
            "fact_id": "FACT-GLOBAL-RANK-30",
            "evidence_ids": ["EVID-RANK-001"],
            "key_metric": "100.0",
            "analytical_intent": "RANKING",
            "visual_spec": {"chart_type": "ranked_bar", "categories": categories, "values": values},
        }),
    }

    budgeted_slides = SpatialCanvasBudgeter.plan_slide_continuations(raw_slide, theme_id="executive_dark")
    # 56-char labels wrap to 2 lines in the 195pt gutter -> 7 items per page -> 5 pages (7+7+7+7+2)
    assert len(budgeted_slides) == 5

    collected_cats = []
    collected_vals = []
    expected_start = 0
    for idx, ps in enumerate(budgeted_slides, start=1):
        _assert_physical_slide_contract(ps, expected_count=len(budgeted_slides))
        part_cats = ps["chart"]["categories"]
        part_vals = ps["chart"]["series"][0]["values"]
        expected_end = expected_start + len(part_cats)
        assert f"Global Ranks #{expected_start + 1} to #{expected_end}" in ps["subtitle"]
        assert ps["content_slice"]["category_slice"] == [expected_start, expected_end]
        collected_cats.extend(part_cats)
        collected_vals.extend(part_vals)
        expected_start = expected_end
        assert ps["semantic_hash"] == raw_slide["semantic_hash"]

    assert collected_cats == categories
    assert collected_vals == values


def test_three_consumer_one_pass_page_boundary_parity(tmp_path: Path):
    """Proves planned physical slide count == Studio pages == PPTX slides == PDF pages across all 4 critical fixtures."""
    headers_b = ["Facility ID", "Turbine Temperature", "Vibration Score", "Compliance Status"]
    rows_b = [[f"FAC-{i+1:03d}", f"{68.0 + i:.1f}°C", f"{0.15 + i*0.01:.2f}g", "Optimal"] for i in range(25)]

    fixture_a = {
        "id": "fix_a",
        "topic_id": "TOPIC-FIX-A",
        "fact_id": "FACT-A",
        "evidence_ids": ["EV-A"],
        "layout": "chart_narrative",
        "title": "Comprehensive Long-Term Strategic Operational Review Analyzing Cross-Functional Interdependencies Across Workforce Demographics and Productivity Dynamics",
        "subtitle": "Cross-Functional Telemetry Analysis",
        "narrative": _build_word_narrative(220),
        "chart": {
            "type": "ranked_bar",
            "categories": ["Region North", "Region South", "Region East", "Region West"],
            "series": [{"name": "Throughput Index", "values": [88.4, 76.2, 91.5, 68.9]}],
            "reference_line": 80.0,
        },
        "semantic_hash": "hash_a",
    }
    fixture_b = {
        "id": "fix_b",
        "topic_id": "TOPIC-FIX-B",
        "fact_id": "FACT-B",
        "evidence_ids": ["EV-B"],
        "layout": "table_detail",
        "title": "Comprehensive High-Precision Turbine Thermal Inspection & Audit",
        "subtitle": "Predictive Diagnostic Asset Health Review",
        "narrative": "Thermal anomaly detected in Facility FAC-018 requiring immediate maintenance window.",
        "image_url": "https://assets.highview.ai/facilities/landscape_plant_01.png",
        "image_shape": "landscape",
        "caption": "Thermal infrared scan of turbine assembly.",
        "table": {"headers": headers_b, "rows": rows_b},
        "semantic_hash": "hash_b",
    }
    fixture_c = {
        "id": "fix_c",
        "topic_id": "TOPIC-FIX-C",
        "fact_id": "FACT-C",
        "evidence_ids": ["EV-C"],
        "layout": "table_detail",
        "title": "Enterprise Departmental Workforce & Financial Audit",
        "subtitle": "Q3 Cross-Functional Ledger",
        "narrative": "Detailed audit of headcount, attendance rates, and operational variance by business unit.",
        "table": {
            "headers": ["Department", "Headcount", "Active FTE", "Total Attendance", "Leave Hours", "Overtime Hours", "Budget $", "Variance %"],
            "rows": [
                ["Engineering", "120", "118", "94.2%", "34h", "12h", "$1,450,000", "+3.2%"],
                ["Product", "45", "44", "91.8%", "18h", "4h", "$620,000", "-1.1%"],
                ["Design", "25", "25", "96.0%", "10h", "2h", "$340,000", "+0.4%"],
                ["Marketing", "35", "33", "88.5%", "22h", "6h", "$480,000", "+2.8%"],
                ["Operations", "80", "78", "92.1%", "28h", "8h", "$950,000", "-0.5%"],
            ],
        },
        "semantic_hash": "hash_c",
    }
    fixture_d = {
        "id": "fix_d",
        "topic_id": "TOPIC-FIX-D",
        "fact_id": "FACT-D",
        "evidence_ids": ["EV-D"],
        "layout": "chart_narrative",
        "title": "Comprehensive Global Business Unit Productivity Ranking",
        "subtitle": "Performance Index Across All 30 Enterprise Nodes",
        "narrative": "Global ranking of all 30 enterprise operational nodes against corporate productivity benchmark.",
        "chart": {
            "type": "ranked_bar",
            "categories": [f"Strategic Business Unit {i+1:02d} — High-Volume Operations Node" for i in range(30)],
            "series": [{"name": "Productivity Index", "values": [float(100.0 - i * 2.5) for i in range(30)]}],
            "reference_line": 75.0,
        },
        "semantic_hash": "hash_d",
    }

    from app.services.presentation.visual.design_tokens import slide_theme_preset
    theme_dict = slide_theme_preset("executive_dark")

    for name, raw_slide in [("Fixture_A", fixture_a), ("Fixture_B", fixture_b), ("Fixture_C", fixture_c), ("Fixture_D", fixture_d)]:
        deck_spec = apply_spatial_slide_budgeting({
            "id": f"deck_{name}",
            "title": f"Track 3.3 {name}",
            "theme": theme_dict,
            "slides": [raw_slide],
            "metadata": {"theme_id": "executive_dark", "title": f"Track 3.3 {name}"},
        })
        planned_count = len(deck_spec["slides"])

        # 1. Studio consumer (resolve_slide)
        studio_pages = 0
        for ps in deck_spec["slides"]:
            plan = resolve_slide(ps, "executive_dark")
            assert plan is not None
            assert len(plan["pages"]) == 1, f"{name}: Studio resolver split a pre-budgeted slide!"
            studio_pages += len(plan["pages"])

        # 2. PPTX consumer
        pptx_path = export_spec_to_pptx(deck_spec)
        prs = Presentation(str(pptx_path))
        pptx_pages = len(prs.slides)

        # 3. PDF consumer
        pdf_path = tmp_path / f"{name}.pdf"
        export_resolved_pdf(deck_spec, pdf_path)
        pdf_pages = _count_pdf_pages(pdf_path)

        assert planned_count == studio_pages == pptx_pages == pdf_pages, (
            f"{name} parity mismatch: planned={planned_count}, studio={studio_pages}, pptx={pptx_pages}, pdf={pdf_pages}"
        )


# ==============================================================================
# 4. CANONICAL 24/24 SEMANTIC IMMUTABILITY AFTER SPATIAL PLANNING
# ==============================================================================

def test_canonical_24_topics_semantic_immutability_after_spatial_planning(tmp_path: Path):
    """Verifies all 24 governed benchmark topics (8 Workforce, 8 Retail, 8 Environmental) preserve 100% semantic immutability AFTER spatial planning."""
    total_topics_checked = 0

    for sid in [9967416, 9967418, 9967419]:
        with get_connection() as conn:
            ctx = capture_dataset_context(conn, sheet_id=sid)
            ws_ev = collect_workspace_evidence(conn, {"sheet_id": sid})

        deck = generate_presentation_deck_spec(
            scope={"sheet_id": sid, "theme_id": "executive_dark"},
            dataset_context=ctx,
            workspace_evidence=ws_ev,
        )
        assert deck["metadata"]["spatial_certified"] is True
        assert deck["metadata"]["contract_mode"] == "EXECUTIVE_TOPIC_BOUND"

        gov_topics = ctx.get("executive_topics") or deck.get("presentation_topics") or []
        assert len(gov_topics) == 8

        # Map physical slides by topic_id
        slides_by_topic: dict[str, list[dict]] = {}
        for ps in deck["slides"]:
            tid = ps.get("topic_id")
            if tid:
                slides_by_topic.setdefault(tid, []).append(ps)

        for gt in gov_topics:
            pt = PresentationTopic.from_executive_topic(gt)
            pre_sem_dict = pt.to_semantic_dict()
            pre_sem_hash = pt.canonical_semantic_hash()

            topic_slides = slides_by_topic.get(pt.topic_id) or []
            assert len(topic_slides) >= 1, f"Missing physical slides for {pt.topic_id}"

            for ps in topic_slides:
                _assert_physical_slide_contract(ps, expected_count=len(topic_slides))
                post_sem_hash = compute_canonical_semantic_hash(ps)
                assert ps["semantic_hash"] == pre_sem_hash == post_sem_hash, f"Semantic hash mutated on {pt.topic_id}"
                assert ps["topic_id"] == pre_sem_dict["topic_id"]
                assert str(ps.get("fact_id") or "") == pre_sem_dict["fact_id"]
                assert sorted(ps.get("evidence_ids") or []) == pre_sem_dict["evidence_ids"]
                assert str(ps.get("key_metric") or "") == pre_sem_dict["metric"]
                assert str(ps.get("unit") or "") == pre_sem_dict["unit"]
                assert str(ps.get("leader") or "") == pre_sem_dict["leader"]
                assert str(ps.get("ranking_direction") or "") == pre_sem_dict["ranking_direction"]
                assert str(ps.get("analytical_intent") or "") == pre_sem_dict["analytical_intent"]
                gt_vspec = gt.get("visual_spec") if isinstance(gt, dict) else gt.visual_spec
                assert ps.get("visual_spec") == gt_vspec

            total_topics_checked += 1

        # Verify 3-consumer page-boundary parity on the full benchmark deck
        planned_n = len(deck["slides"])
        studio_n = sum(len(resolve_slide(s, "executive_dark")["pages"]) for s in deck["slides"])
        pptx_path = export_spec_to_pptx(deck)
        pptx_n = len(Presentation(str(pptx_path)).slides)
        pdf_path = tmp_path / f"bench_{sid}.pdf"
        export_resolved_pdf(deck, pdf_path)
        pdf_n = _count_pdf_pages(pdf_path)
        assert planned_n == studio_n == pptx_n == pdf_n

    assert total_topics_checked == 24


# ==============================================================================
# 5. COMPLETE STRESS COVERAGE MATRIX
# ==============================================================================

@pytest.mark.parametrize("row_count", [8, 15, 25, 50, 100])
def test_stress_matrix_table_row_counts(row_count: int):
    headers = ["Entity ID", "Primary Metric", "Benchmark Delta", "Audit Status"]
    rows = [[f"ENT-{i+1:03d}", f"{92.4 - i*0.2:.1f}%", f"{+1.5 - i*0.05:+.2f}%", "Verified"] for i in range(row_count)]
    slide = {
        "id": f"stress_rows_{row_count}",
        "topic_id": f"TOPIC-ROWS-{row_count}",
        "fact_id": "FACT-ROWS",
        "evidence_ids": ["EV-1"],
        "layout": "table_detail",
        "title": f"Enterprise Audit Ledger ({row_count} Rows)",
        "subtitle": "Deterministic Row Pagination Stress Test",
        "narrative": "Full row conservation stress test.",
        "table": {"headers": headers, "rows": rows},
        "semantic_hash": "hash_rows",
    }
    planned = SpatialCanvasBudgeter.plan_slide_continuations(slide)
    collected = []
    for ps in planned:
        _assert_physical_slide_contract(ps, expected_count=len(planned))
        assert ps["table"]["headers"] == headers
        collected.extend(ps["table"]["rows"])
    assert collected == rows


@pytest.mark.parametrize("col_count", [4, 8, 12, 20])
def test_stress_matrix_table_column_counts(col_count: int):
    headers = ["Entity ID"] + [f"Metric_{c:02d}" for c in range(1, col_count)]
    rows = [[f"NODE-{r+1:02d}"] + [f"{(r+1)*10 + c}" for c in range(1, col_count)] for r in range(6)]
    slide = {
        "id": f"stress_cols_{col_count}",
        "topic_id": f"TOPIC-COLS-{col_count}",
        "fact_id": "FACT-COLS",
        "evidence_ids": ["EV-1"],
        "layout": "table_detail",
        "title": f"Wide Table Column Stress Test ({col_count} Columns)",
        "subtitle": "Identity Column Repetition Verification",
        "narrative": "Column group pagination stress test.",
        "table": {"headers": headers, "rows": rows},
        "semantic_hash": "hash_cols",
    }
    planned = SpatialCanvasBudgeter.plan_slide_continuations(slide)
    expected_pages = 1 if col_count <= 4 else ((col_count - 1 + 2) // 3)
    assert len(planned) == expected_pages
    collected_val_cols = []
    for ps in planned:
        _assert_physical_slide_contract(ps, expected_count=expected_pages)
        assert ps["table"]["headers"][0] == "Entity ID"
        collected_val_cols.extend(ps["table"]["headers"][1:])
    assert collected_val_cols == headers[1:]


@pytest.mark.parametrize("cat_count", [5, 10, 15, 25, 50])
@pytest.mark.parametrize("long_labels", [False, True])
def test_stress_matrix_chart_category_counts(cat_count: int, long_labels: bool):
    cats = [
        (f"Global Enterprise Operating Unit & Logistics Distribution Center #{i+1:02d}" if long_labels else f"Unit {i+1:02d}")
        for i in range(cat_count)
    ]
    vals = [float(200.0 - i * 1.5) for i in range(cat_count)]
    slide = {
        "id": f"stress_cats_{cat_count}_{long_labels}",
        "topic_id": f"TOPIC-CATS-{cat_count}",
        "fact_id": "FACT-CATS",
        "evidence_ids": ["EV-1"],
        "layout": "chart_narrative",
        "title": f"Category Scaling Stress Test ({cat_count} Items)",
        "subtitle": "Adaptive Plot Capacity Check",
        "narrative": "Category ranking stress test.",
        "chart": {
            "type": "ranked_bar",
            "categories": cats,
            "series": [{"name": "Score", "values": vals}],
            "reference_line": 150.0,
        },
        "semantic_hash": "hash_cats",
    }
    planned = SpatialCanvasBudgeter.plan_slide_continuations(slide)
    collected_cats = []
    collected_vals = []
    for ps in planned:
        _assert_physical_slide_contract(ps, expected_count=len(planned))
        collected_cats.extend(ps["chart"]["categories"])
        collected_vals.extend(ps["chart"]["series"][0]["values"])
    assert collected_cats == cats
    assert collected_vals == vals


@pytest.mark.parametrize("image_shape", ["landscape", "portrait", "square", "very_wide", "very_tall"])
def test_stress_matrix_image_shapes(image_shape: str, tmp_path: Path):
    slide = {
        "id": f"stress_img_{image_shape}",
        "topic_id": f"TOPIC-IMG-{image_shape.upper()}",
        "fact_id": "FACT-IMG",
        "evidence_ids": ["EV-IMG"],
        "layout": "comparison_split",
        "title": f"Visual Asset Aspect-Ratio Stress Test ({image_shape})",
        "subtitle": "Zero Distortion & Zero Collision Check",
        "narrative": "Visual inspection confirms aspect-ratio-preserving containment within right column.",
        "image_url": f"https://assets.highview.ai/inspections/{image_shape}_sensor.png",
        "image_shape": image_shape,
        "caption": f"Certified {image_shape} diagnostic visual frame.",
        "semantic_hash": f"hash_img_{image_shape}",
    }
    from app.services.presentation.visual.design_tokens import slide_theme_preset
    deck = apply_spatial_slide_budgeting({
        "id": f"deck_img_{image_shape}",
        "title": f"Image {image_shape}",
        "theme": slide_theme_preset("executive_dark"),
        "slides": [slide],
        "metadata": {"theme_id": "executive_dark"},
    })
    assert len(deck["slides"]) == 1
    _assert_physical_slide_contract(deck["slides"][0], expected_count=1)
    pptx_path = export_spec_to_pptx(deck)
    assert pptx_path.exists()
    pdf_path = tmp_path / f"img_{image_shape}.pdf"
    export_resolved_pdf(deck, pdf_path)
    assert pdf_path.exists()


@pytest.mark.parametrize("narrative_words", [50, 100, 200, 400, 800])
def test_stress_matrix_narrative_lengths(narrative_words: int):
    text = _build_word_narrative(narrative_words)
    slide = {
        "id": f"stress_narr_{narrative_words}",
        "topic_id": f"TOPIC-NARR-{narrative_words}",
        "fact_id": "FACT-NARR",
        "evidence_ids": ["EV-NARR"],
        "layout": "comparison_split",
        "title": f"Narrative Length Stress Test ({narrative_words} Words)",
        "subtitle": "Deterministic Multi-Slide Narrative Flow",
        "narrative": text,
        "semantic_hash": f"hash_narr_{narrative_words}",
    }
    planned = SpatialCanvasBudgeter.plan_slide_continuations(slide)
    reconstructed = [w for ps in planned for w in (ps.get("narrative") or "").split()]
    assert reconstructed == text.split()
    for ps in planned:
        _assert_physical_slide_contract(ps, expected_count=len(planned))


# ==============================================================================
# 6. ZERO SLIDELAYOUTERROR ESCAPE & TYPED SPATIALPLANNINGERROR
# ==============================================================================

def test_zero_slide_layout_error_escape_and_typed_spatial_planning_error():
    """Proves unsolvable geometry raises typed SpatialPlanningError and 0 SlideLayoutError escapes."""
    unsolvable_cases = [
        # 1. Unsupported layout
        {"title": "Bad Layout", "layout": "hologram_3d_canvas"},
        # 2. Empty table
        {"title": "Empty Table", "layout": "table_detail", "table": {"headers": ["A", "B"], "rows": []}},
        # 3. Extra table cells beyond headers
        {"title": "Extra Cells", "layout": "table_detail", "table": {"headers": ["A", "B"], "rows": [["1", "2", "3"]]}},
        # 4. Single atomic table row taller than the entire slide canvas (>350pt)
        {
            "title": "Impossible Atomic Cell",
            "layout": "table_detail",
            "table": {"headers": ["ID", "Notes"], "rows": [["1", "WORD " * 600]]},
        },
        # 5. Invalid image shape
        {
            "title": "Invalid Image Shape",
            "layout": "comparison_split",
            "image_url": "https://assets.highview.ai/x.png",
            "image_shape": "hypercube_4d",
        },
        # 6. Excessive metrics (>8) on a single slide
        {
            "title": "Too Many Metrics",
            "layout": "key_metrics",
            "metrics": [{"label": f"M{i}", "value": str(i)} for i in range(10)],
        },
        # 7. Oversized footer source_label
        {
            "title": "Huge Footer Source Label",
            "layout": "comparison_split",
            "narrative": "Valid body.",
            "source_label": "Source: " + ("Long citation segment " * 30),
        },
    ]

    for case in unsolvable_cases:
        with pytest.raises(SpatialPlanningError):
            SpatialCanvasBudgeter.plan_slide_continuations(case)


# ==============================================================================
# 7. SPATIAL PERFORMANCE & ZERO MODEL CALLS
# ==============================================================================

def test_spatial_planning_performance_and_zero_model_calls():
    """Benchmarks deterministic spatial planning latency (mean, P50, P95) with 0 model calls."""
    import app.services.presentation.ai_deck_planner as ai_planner
    import app.services.presentation.orchestrator.granite_orchestrator as granite_orch

    llm_calls = 0

    def fail_if_llm_called(*args, **kwargs):
        nonlocal llm_calls
        llm_calls += 1
        raise RuntimeError("LLM invoked during spatial planning!")

    simple_slide = {
        "id": "perf_simple",
        "topic_id": "TOPIC-SIMPLE",
        "layout": "chart_narrative",
        "title": "Simple Quarterly Throughput",
        "subtitle": "4 Regions",
        "narrative": "Concise executive takeaway.",
        "chart": {
            "type": "ranked_bar",
            "categories": ["North", "South", "East", "West"],
            "series": [{"name": "Throughput", "values": [91.0, 88.0, 84.0, 79.0]}],
        },
        "semantic_hash": "h_simple",
    }

    paginated_slide = {
        "id": "perf_paginated",
        "topic_id": "TOPIC-PAGINATED",
        "layout": "chart_narrative",
        "title": "Paginated 30-Unit Long-Label Ranking",
        "subtitle": "30 Enterprise Nodes",
        "narrative": _build_word_narrative(200),
        "chart": {
            "type": "ranked_bar",
            "categories": [f"Strategic Business Unit {i+1:02d} — High-Volume Operations Node" for i in range(30)],
            "series": [{"name": "Index", "values": [float(100 - i) for i in range(30)]}],
            "reference_line": 75.0,
        },
        "semantic_hash": "h_paginated",
    }

    canonical_8_deck = {
        "id": "perf_8_deck",
        "title": "Canonical 8-Topic Deck",
        "theme": "executive_dark",
        "slides": [
            {"id": "cover", "layout": "title_cover", "title": "Executive Briefing", "subtitle": "8 Governed Topics"},
            *[{**simple_slide, "id": f"s_{i}", "topic_id": f"TOPIC-{i}"} for i in range(8)],
        ],
        "metadata": {"theme_id": "executive_dark"},
    }

    stress_deck = {
        "id": "perf_stress_deck",
        "title": "Worst-Case Stress Deck",
        "theme": "executive_dark",
        "slides": [
            {"id": "cover", "layout": "title_cover", "title": "Stress Briefing", "subtitle": "Full Matrix"},
            paginated_slide,
            {
                "id": "stress_100_rows",
                "topic_id": "TOPIC-100-ROWS",
                "layout": "table_detail",
                "title": "100-Row Audit Table",
                "subtitle": "Full Conservation",
                "narrative": "100-row table.",
                "table": {
                    "headers": ["ID", "Metric", "Delta", "Status"],
                    "rows": [[f"ID-{i}", str(i), "+1%", "OK"] for i in range(100)],
                },
                "semantic_hash": "h_100",
            },
            {
                "id": "stress_20_cols",
                "topic_id": "TOPIC-20-COLS",
                "layout": "table_detail",
                "title": "20-Column Wide Table",
                "subtitle": "Identity Safety",
                "narrative": "20-column table.",
                "table": {
                    "headers": ["Department"] + [f"Col_{c}" for c in range(1, 20)],
                    "rows": [[f"Dept_{r}"] + [str(c) for c in range(1, 20)] for r in range(8)],
                },
                "semantic_hash": "h_20c",
            },
            {
                "id": "stress_800_words",
                "topic_id": "TOPIC-800-WORDS",
                "layout": "comparison_split",
                "title": "800-Word Visible Narrative",
                "subtitle": "Narrative Flow",
                "narrative": _build_word_narrative(800),
                "semantic_hash": "h_800w",
            },
        ],
        "metadata": {"theme_id": "executive_dark"},
    }

    with patch.object(ai_planner, "plan_deck_with_ai", side_effect=fail_if_llm_called):
        with patch.object(granite_orch, "PresentationExecutionOrchestrator", side_effect=fail_if_llm_called):
            for workload_name, fn in [
                ("single_simple_topic", lambda: SpatialCanvasBudgeter.plan_slide_continuations(simple_slide)),
                ("single_paginated_topic", lambda: SpatialCanvasBudgeter.plan_slide_continuations(paginated_slide)),
                ("canonical_8_topic_deck", lambda: apply_spatial_slide_budgeting(canonical_8_deck)),
                ("worst_case_stress_deck", lambda: apply_spatial_slide_budgeting(stress_deck)),
            ]:
                timings_ms = []
                for _ in range(15):
                    t0 = time.perf_counter()
                    res = fn()
                    timings_ms.append((time.perf_counter() - t0) * 1000.0)
                mean_ms = statistics.mean(timings_ms)
                p50_ms = statistics.median(timings_ms)
                p95_ms = sorted(timings_ms)[int(len(timings_ms) * 0.95) - 1]
                assert mean_ms < 250.0, f"{workload_name} too slow: {mean_ms:.2f}ms"
                assert p50_ms < 250.0
                assert p95_ms < 350.0

    assert llm_calls == 0
