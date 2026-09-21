import json
import zipfile
import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from fastapi.testclient import TestClient

from app.main import app
from app.db.database import get_connection
from app.services.presentation_service import (
    THEMES,
    PresentationJobManager,
    capture_dataset_context,
    extract_presentation_charts,
    synthesize_presentation_charts,
    generate_presentation_deck_spec,
    regenerate_single_slide,
    execute_presentation_pipeline_async,
    collect_workspace_evidence,
    verify_presentation_claims
)
from app.services.report_generator import export_spec_to_pptx, ChartExportError
from app.services.shared_evidence_package import build_shared_evidence_package, generate_coverage_manifest
from app.services.presentation_quality_auditor import PresentationQualityAuditor



def seed_test_sales_sheet(client: TestClient | None = None) -> int:
    if client is None:
        client = TestClient(app)
    csv_lines = ["Store,Date,Weekly_Sales,Holiday_Flag"]
    for i in range(1, 101):
        store = f"Store {((i % 15) + 1)}"
        date = f"2023-{(i % 12) + 1:02d}-01"
        sales = 15000.0 + (i * 250.0)
        flag = 1 if (i % 5 == 0) else 0
        csv_lines.append(f"{store},{date},{sales},{flag}")
    csv_bytes = "\n".join(csv_lines).encode("utf-8")
    resp = client.post("/api/upload/file", files={"file": ("walmart_sales.csv", csv_bytes, "text/csv")})
    assert resp.status_code == 200
    sheets_resp = client.get("/api/sheets")
    assert sheets_resp.status_code == 200
    sheets = sheets_resp.json()["sheets"]
    return sheets[0]["id"]


def test_theme_registry():
    assert "executive_dark" in THEMES
    assert "clean_light" in THEMES
    assert "corporate_navy" in THEMES
    assert "emerald_slate" in THEMES
    assert "bold_signal" in THEMES
    assert "electric_studio" in THEMES
    assert "creative_voltage" in THEMES
    assert "dark_botanical" in THEMES
    assert "swiss_modern" in THEMES
    assert "notebook_tabs" in THEMES

    for key, t in THEMES.items():
        assert "bg_color" in t
        assert "card_bg" in t
        assert "brand_color" in t
        assert "chart_palette" in t
        assert len(t["chart_palette"]) >= 5


@pytest.fixture(autouse=True)
def fast_tests(monkeypatch):
    monkeypatch.setattr("app.services.presentation_service._call_ai_presentation_enrichment", lambda **kwargs: None)


def test_capture_dataset_context_and_charts():
    sheet_id = seed_test_sales_sheet()
    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)

    assert ctx["target_sheet"]["id"] == sheet_id
    assert "Sales" in ctx["domain"] or "Commercial" in ctx["domain"]
    assert len(ctx["records"]) == 100
    assert ctx["snapshot_hash"] is not None

    charts = extract_presentation_charts(ctx["visuals"], ctx["domain"])
    assert len(charts) > 0

    for chart in charts:
        if chart.get("type") in ("bar", "horizontal_bar"):
            assert len(chart.get("categories", [])) <= 10


def test_deck_spec_generation_sales_domain():
    sheet_id = seed_test_sales_sheet()
    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)

    scope = {
        "objective": "Quarterly Sales & Store Distribution Review",
        "audience": "Executive Vice Presidents",
        "target_length": 6,
        "theme_id": "corporate_navy",
        "instructions": "Emphasize store performance rankings."
    }

    spec = generate_presentation_deck_spec(scope, ctx)

    assert spec["spec_version"] == "2.0"
    assert spec["theme"]["id"] == "corporate_navy"
    assert spec["metadata"]["domain"] == ctx["domain"]
    assert len(spec["slides"]) >= 5

    layouts = [s["layout"] for s in spec["slides"]]
    assert "title_hero" in layouts
    assert "kpi_summary" in layouts
    assert "chart_narrative" in layouts
    assert "comparison_split" in layouts
    assert "table_detail" in layouts

    all_text = " ".join([
        s.get("title", "") + " " + s.get("narrative", "") + " " + " ".join(s.get("bullets", []))
        for s in spec["slides"]
    ]).lower()
    assert "burnout" not in all_text
    assert "flight risk" not in all_text

    for s in spec["slides"]:
        assert len(s.get("speaker_notes", "")) > 10
        assert len(s.get("evidence_sources", [])) > 0


def test_export_spec_to_pptx_with_native_charts():
    sheet_id = seed_test_sales_sheet()
    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)

    spec = generate_presentation_deck_spec({
        "objective": "Store Throughput Review",
        "theme_id": "emerald_slate"
    }, ctx)

    pptx_path = export_spec_to_pptx(spec)
    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 10000

    prs = Presentation(str(pptx_path))
    assert abs(prs.slide_width.inches - 13.333) < 0.01
    assert abs(prs.slide_height.inches - 7.5) < 0.01
    assert len(prs.slides) == len(spec["slides"])

    has_at_least_one_chart = False
    for slide in prs.slides:
        if slide.has_notes_slide:
            assert len(slide.notes_slide.notes_text_frame.text) > 0
        chart_count = sum(1 for s in slide.shapes if s.has_chart)
        if chart_count > 0:
            has_at_least_one_chart = True
            for s in slide.shapes:
                if s.has_chart:
                    assert len(s.chart.series) > 0

    assert has_at_least_one_chart is True


def test_regenerate_single_slide():
    sheet_id = seed_test_sales_sheet()
    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)

    spec = generate_presentation_deck_spec({"objective": "General Overview"}, ctx)
    target_slide = spec["slides"][1]
    slide_id = target_slide["id"]

    prompt = "Focus sharply on holiday promotional uplift and inventory prep."
    updated_deck = regenerate_single_slide(spec, slide_id, prompt)

    updated_slide = next(s for s in updated_deck["slides"] if s["id"] == slide_id)
    assert len(updated_slide["bullets"]) >= 2
    assert prompt.lower() in updated_slide["speaker_notes"].lower() or len(updated_slide["speaker_notes"]) > 10


def test_presentation_job_manager_cancellation():
    manager = PresentationJobManager()
    job_id = manager.create_job({"objective": "Test Job"})

    job = manager.get_job(job_id)
    assert job["status"] in ("in_progress", "pending")
    assert job["stage"] in ("reviewing_coverage", "collecting_findings", "queued")

    manager.update_stage(job_id, "planning_outline", "Structuring presentation slides", 35)
    job = manager.get_job(job_id)
    assert job["stage"] == "planning_outline"
    assert job["progress_pct"] == 35

    assert manager.is_cancelled(job_id) is False
    manager.cancel_job(job_id)
    assert manager.is_cancelled(job_id) is True

    job = manager.get_job(job_id)
    assert job["status"] == "cancelled"


def test_presentations_api_endpoints():
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)

    # 1. Themes
    res = client.get("/api/presentations/themes")
    assert res.status_code == 200
    themes = res.json()["themes"]
    assert len(themes) >= 4

    # 2. Scope preview
    prev_res = client.post("/api/presentations/scope-preview", json={
        "scope_type": "workspace"
    })
    assert prev_res.status_code == 200
    prev_data = prev_res.json()
    assert prev_data["eligible"] is True
    assert prev_data["total_included_sheets"] >= 1
    assert "reporting_period_summary" in prev_data

    # 3. Start generation job with scope
    res = client.post("/api/presentations/generate", json={
        "objective": "API Integration Test",
        "theme_id": "clean_light",
        "scope_type": "workspace"
    })
    assert res.status_code == 200
    job_id = res.json()["job_id"]
    assert job_id is not None
    assert res.json()["stage"] == "reviewing_coverage"

    # 4. Poll job
    res = client.get(f"/api/presentations/jobs/{job_id}")
    assert res.status_code == 200
    assert "stage" in res.json()

    # 5. Cancel job
    res = client.post(f"/api/presentations/jobs/{job_id}/cancel")
    assert res.status_code == 200
    assert res.json()["status"] == "cancelled"


def test_ready_job_attaches_deck():
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)
    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)
    spec = generate_presentation_deck_spec({"objective": "Test Ready Deck"}, ctx)
    deck_id = spec["id"]
    with get_connection() as conn:
        conn.execute(
            "INSERT INTO presentation_decks (id, title, dataset_id, sheet_id, theme_id, spec_json, pptx_filename, created_at, updated_at) "
            "VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)",
            (deck_id, spec["metadata"]["title"], None, sheet_id, spec["metadata"]["theme_id"], json.dumps(spec), "test.pptx")
        )
        conn.commit()

    from app.routers.presentations import job_manager as router_job_mgr
    job_id = router_job_mgr.create_job({"objective": "Test Ready Deck"})
    router_job_mgr.update_stage(job_id, "ready", "Presentation ready to review", 100, deck_id=deck_id)

    res = client.get(f"/api/presentations/jobs/{job_id}")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "ready"
    assert "deck" in data
    assert data["deck"]["id"] == deck_id
    assert len(data["deck"]["slides"]) >= 4

    # Test evidence inspection endpoint
    ev_res = client.get(f"/api/presentations/decks/{deck_id}/evidence")
    assert ev_res.status_code == 200
    ev_data = ev_res.json()
    assert ev_data["deck_id"] == deck_id
    assert len(ev_data["evidence_ledger"]) > 0

    # Test revalidation endpoint
    reval_res = client.post("/api/presentations/revalidate", json={"deck_spec": spec})
    assert reval_res.status_code == 200
    reval_data = reval_res.json()
    assert reval_data["status"] == "PASSED"
    assert reval_data["total_metrics_checked"] > 0


def test_detect_sheet_date_range_and_partial_year():
    from app.services.presentation_service import detect_sheet_date_range

    # 1. Partial year (e.g. 5 months)
    records_partial = [
        {"Date": f"2023-0{m}-15", "Sales": 1000 * m}
        for m in range(1, 6)
    ]
    res_partial = detect_sheet_date_range(records_partial, ["Date", "Sales"])
    assert res_partial["has_date"] is True
    assert res_partial["is_partial_year"] is True
    assert "Partial Year" in res_partial["period_label"]

    # 2. Multi-year (> 2 years)
    records_multi = [
        {"Date": "2020-01-01", "Sales": 100},
        {"Date": "2021-01-01", "Sales": 200},
        {"Date": "2023-01-01", "Sales": 300}
    ]
    res_multi = detect_sheet_date_range(records_multi, ["Date", "Sales"])
    assert res_multi["has_date"] is True
    assert res_multi["is_partial_year"] is False
    assert "Years" in res_multi["period_label"]


def test_workspace_evidence_collection_and_eight_slides():
    from app.services.presentation_service import (
        collect_workspace_evidence,
        generate_presentation_deck_spec,
        verify_presentation_claims
    )
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)

    scope = {
        "scope_type": "workspace",
        "objective": "Consolidated Board Review",
        "audience": "Board of Directors",
        "theme_id": "corporate_navy"
    }

    with get_connection() as conn:
        evidence = collect_workspace_evidence(conn, scope)

    assert evidence["total_records"] >= 100
    assert len(evidence["evidence_ledger"]) == 8
    assert evidence["snapshot_hash"] is not None

    evidence_ids = [e["evidence_id"] for e in evidence["evidence_ledger"]]
    assert "EVID-EXEC-01" in evidence_ids
    assert "EVID-KPI-01" in evidence_ids
    assert "EVID-STRENGTH-01" in evidence_ids
    assert "EVID-HEADWIND-01" in evidence_ids
    assert "EVID-REL-01" in evidence_ids
    assert "EVID-LESSON-01" in evidence_ids
    assert "EVID-REC-01" in evidence_ids
    assert "EVID-GOV-01" in evidence_ids

    # Generate deck (flexible multi-slide section count: scales from 8+ up to 14-16 based on evidence complexity)
    deck_spec = generate_presentation_deck_spec(scope, evidence["primary_ctx"], workspace_evidence=evidence)
    assert len(deck_spec["slides"]) >= 8
    assert deck_spec["metadata"]["validation_summary"]["status"] == "PASSED"

    # Check that Action Plan slide contains structured proposals with unassigned owner roles
    slide_7 = next(s for s in deck_spec["slides"] if s.get("stable_slide_id") == "slide_action_plan" or s.get("layout") == "action_plan")
    assert slide_7["category"] in ("STRATEGIC PROPOSALS", "STRATEGIC ROADMAP")
    assert "structured_proposals" in slide_7
    for prop in slide_7["structured_proposals"]:
        assert "motivating_finding" in prop
        assert "proposed_response" in prop
        assert "priority" in prop
        assert "owner_role" in prop
        assert "Unassigned" in prop["owner_role"]
        assert "success_metric" in prop
        assert "dependencies" in prop

    # Check that speaker notes contain Presenter Briefing (Board Scrutiny)
    for s in deck_spec["slides"]:
        assert "=== PRESENTER BRIEFING (BOARD SCRUTINY) ===" in s["speaker_notes"]
        assert "• How Calculated:" in s["speaker_notes"]
        assert "• What it Establishes:" in s["speaker_notes"]
        assert "• What it Does NOT Establish:" in s["speaker_notes"]
        assert "• Anticipated Board Q&A:" in s["speaker_notes"]

    # Verify claim verifier
    ver_res = verify_presentation_claims(deck_spec, evidence["evidence_ledger"])
    assert ver_res["status"] == "PASSED"
    assert ver_res["passed_verification"] > 0
    assert ver_res["discrepancies_flagged"] == 0


def test_native_chart_pie_vs_donut_mapping():
    """Verify pie and donut charts are mapped distinctly to XL_CHART_TYPE.PIE and XL_CHART_TYPE.DOUGHNUT."""
    pie_spec = {
        "id": "deck_pie_test",
        "theme": {"id": "executive_dark"},
        "slides": [
            {
                "id": "s_pie",
                "layout": "chart_narrative",
                "title": "Headcount by Region (Pie)",
                "subtitle": "Part-to-whole distribution",
                "narrative": "Regional allocation across primary business units.",
                "bullets": ["North region comprises 45% of headcount."],
                "metrics": [{"label": "Total Staff", "value": "1,000", "subtext": "Complete cohort"}],
                "chart": {
                    "chart_type": "pie",
                    "title": "Headcount by Region",
                    "categories": ["North", "South", "East", "West"],
                    "series": [{"name": "Share", "values": [450.0, 250.0, 180.0, 120.0]}]
                },
                "evidence_sources": ["Workforce DB"]
            }
        ]
    }
    pie_path = export_spec_to_pptx(pie_spec)
    prs_pie = Presentation(str(pie_path))
    chart_shape = next(s for s in prs_pie.slides[0].shapes if s.has_chart)
    assert chart_shape.chart.chart_type == XL_CHART_TYPE.PIE

    donut_spec = {
        "id": "deck_donut_test",
        "theme": {"id": "executive_dark"},
        "slides": [
            {
                "id": "s_donut",
                "layout": "chart_narrative",
                "title": "Headcount by Region (Donut)",
                "subtitle": "Ring distribution",
                "narrative": "Regional allocation across primary business units.",
                "bullets": ["North region comprises 45% of headcount."],
                "metrics": [{"label": "Total Staff", "value": "1,000", "subtext": "Complete cohort"}],
                "chart": {
                    "chart_type": "donut",
                    "title": "Headcount by Region",
                    "categories": ["North", "South", "East", "West"],
                    "series": [{"name": "Share", "values": [450.0, 250.0, 180.0, 120.0]}]
                },
                "evidence_sources": ["Workforce DB"]
            }
        ]
    }
    donut_path = export_spec_to_pptx(donut_spec)
    prs_donut = Presentation(str(donut_path))
    chart_shape = next(s for s in prs_donut.slides[0].shapes if s.has_chart)
    assert chart_shape.chart.chart_type == XL_CHART_TYPE.DOUGHNUT


def test_native_chart_multi_series_pptx_and_zip_parts():
    """Verify multi-series column chart exports with genuine embedded Excel and XML parts."""
    multi_spec = {
        "id": "deck_multiseries_test",
        "theme": {"id": "executive_dark"},
        "slides": [
            {
                "id": "s_multi",
                "layout": "chart_narrative",
                "title": "Comparative Department Diagnostics",
                "subtitle": "Performance vs Absent Days",
                "narrative": "Comparing average performance rating and absenteeism across 4 departments.",
                "bullets": ["Engineering achieved highest performance score."],
                "metrics": [{"label": "Departments", "value": "4", "subtext": "Evaluated units"}],
                "chart": {
                    "chart_type": "column",
                    "title": "Performance vs Absent Days",
                    "categories": ["Engineering", "Sales", "Support", "Marketing"],
                    "series": [
                        {"name": "Performance Rating", "values": [4.5, 3.8, 4.1, 3.9]},
                        {"name": "Absent Days", "values": [2.1, 4.5, 3.2, 5.0]}
                    ]
                },
                "evidence_sources": ["Cross-Sheet Join"]
            }
        ]
    }
    pptx_path = export_spec_to_pptx(multi_spec)
    prs = Presentation(str(pptx_path))
    chart_shape = next(s for s in prs.slides[0].shapes if s.has_chart)
    chart = chart_shape.chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    assert len(chart.series) == 2
    assert chart.series[0].name == "Performance Rating"
    assert chart.series[1].name == "Absent Days"

    # Inspect zip package parts to verify embedded Excel workbook
    with zipfile.ZipFile(str(pptx_path), "r") as z:
        names = z.namelist()
        assert any(n.startswith("ppt/charts/chart") and n.endswith(".xml") for n in names)
        assert any("embeddings" in n and n.endswith(".xlsx") for n in names)


def test_chart_export_validation_fail_fast():
    """Verify PowerPoint exporter fails fast on invalid chart specifications rather than concealing errors."""
    base_slide = {
        "id": "s_err",
        "layout": "chart_narrative",
        "title": "Test Error",
        "subtitle": "Validation Test",
        "narrative": "Validation check.",
        "bullets": ["Bullet 1"],
        "evidence_sources": ["Test Source"]
    }

    # 1. Missing chart on chart_narrative layout
    with pytest.raises(ChartExportError, match="contains no chart specification"):
        export_spec_to_pptx({
            "id": "deck_err_1",
            "slides": [{**base_slide, "chart": None}]
        })

    # 2. Empty categories
    with pytest.raises(ChartExportError, match="empty categories"):
        export_spec_to_pptx({
            "id": "deck_err_2",
            "slides": [{
                **base_slide,
                "chart": {
                    "chart_type": "column",
                    "categories": [],
                    "series": [{"name": "S1", "values": [10.0]}]
                }
            }]
        })

    # 3. Mismatched series/categories lengths
    with pytest.raises(ChartExportError, match="values length .* does not match categories length"):
        export_spec_to_pptx({
            "id": "deck_err_3",
            "slides": [{
                **base_slide,
                "chart": {
                    "chart_type": "line",
                    "categories": ["A", "B", "C"],
                    "series": [{"name": "S1", "values": [10.0, 20.0]}]
                }
            }]
        })

    # 4. Non-finite values
    with pytest.raises(ChartExportError, match="non-finite or null value"):
        export_spec_to_pptx({
            "id": "deck_err_4",
            "slides": [{
                **base_slide,
                "chart": {
                    "chart_type": "line",
                    "categories": ["A", "B"],
                    "series": [{"name": "S1", "values": [10.0, float("nan")]}]
                }
            }]
        })


def test_part_to_whole_pie_other_aggregation():
    """Verify category count > 6 aggregates remainder into 'Other' preserving 100% total population."""
    slices = [
        {"label": f"Cat_{i}", "count": 10.0 * (10 - i)}
        for i in range(10)
    ]
    total_expected = sum(s["count"] for s in slices)
    visuals = [{
        "id": "vis_donut_test",
        "chart_type": "donut",
        "title": "Category Distribution",
        "donut_data": {
            "category_col": "Department",
            "total": total_expected,
            "slices": slices
        }
    }]
    charts = extract_presentation_charts(visuals, "Operations")
    assert len(charts) == 1
    c = charts[0]
    # Slices > 6 should yield top 5 + "Other" = 6 categories
    assert len(c["categories"]) == 6
    assert c["categories"][-1] == "Other"
    values = c["series"][0]["values"]
    assert len(values) == 6
    assert abs(sum(values) - total_expected) < 0.01
    assert "aggregated into 'Other'" in c["aggregation_disclosure"]


def test_synthesis_and_charts_on_wide_attendance_matrix():
    """Verify wide attendance sheets (Date + Person_0...Person_N) synthesize charts for all analytical slides."""
    client = TestClient(app)
    # Generate CSV with Date and 10 Person columns
    csv_lines = ["Date," + ",".join(f"Person_{i}" for i in range(10))]
    for d in range(1, 31):
        row = [f"2023-01-{d:02d}"] + [f"{7.5 + (d % 3) * 0.5 + (i % 2) * 0.2:.1f}" for i in range(10)]
        csv_lines.append(",".join(row))
    csv_bytes = "\n".join(csv_lines).encode("utf-8")
    resp = client.post("/api/upload/file", files={"file": ("attendance_matrix.csv", csv_bytes, "text/csv")})
    assert resp.status_code == 200

    sheets_resp = client.get("/api/sheets")
    assert sheets_resp.status_code == 200
    dataset_id = resp.json().get("dataset_id")
    sheet = next(
        s for s in sheets_resp.json()["sheets"]
        if (dataset_id and s.get("dataset_id") == dataset_id)
        or "attendance_matrix" in s.get("original_name", "")
        or "attendance_matrix" in s.get("name", "")
    )
    sheet_id = sheet["id"]

    with get_connection() as conn:
        ctx = capture_dataset_context(conn, sheet_id=sheet_id)

    scope = {
        "scope_type": "sheet",
        "sheet_id": sheet_id,
        "objective": "Attendance Matrix Executive Review",
        "theme_id": "executive_dark"
    }
    deck_spec = generate_presentation_deck_spec(scope, ctx)

    # Slides 3, 4, 5 must have chart_narrative and populated valid chart specs
    s3 = deck_spec["slides"][2]
    s4 = deck_spec["slides"][3]
    s5 = deck_spec["slides"][4]

    assert s3["layout"] == "chart_narrative"
    assert s3["chart"] is not None
    assert s3["chart"]["chart_type"] == "line"
    assert len(s3["chart"]["categories"]) > 0

    assert s4["layout"] == "chart_narrative"
    assert s4["chart"] is not None
    assert s4["chart"]["chart_type"] in ("bar", "column")
    assert len(s4["chart"]["categories"]) > 0

    assert s5["layout"] == "chart_narrative"
    assert s5["chart"] is not None
    assert len(s5["chart"]["categories"]) > 0

    # Test that exporting to PPTX succeeds and shapes have native charts
    pptx_path = export_spec_to_pptx(deck_spec)
    assert pptx_path.exists()

    prs = Presentation(str(pptx_path))
    # Check that slides 3, 4, 5 each contain at least one native chart shape
    for slide_idx in (2, 3, 4):
        slide = prs.slides[slide_idx]
        chart_shapes = [s for s in slide.shapes if s.has_chart]
        assert len(chart_shapes) >= 1, f"Slide {slide_idx+1} is missing native chart shape!"


def test_shared_evidence_package_and_coverage_manifest():
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)
    scope = {"scope_type": "workspace", "objective": "Coverage Accounting Test"}
    with get_connection() as conn:
        pkg = build_shared_evidence_package(conn, scope)

    assert pkg["snapshot_hash"] is not None
    assert len(pkg["snapshot_hash"]) >= 12
    assert len(pkg["candidate_findings"]) >= 8

    evidence_ids = [f.get("evidence_id") for f in pkg["candidate_findings"] if f.get("evidence_id")]
    for canonical_id in [
        "EVID-EXEC-01", "EVID-KPI-01", "EVID-STRENGTH-01", "EVID-HEADWIND-01",
        "EVID-REL-01", "EVID-LESSON-01", "EVID-REC-01", "EVID-GOV-01"
    ]:
        assert canonical_id in evidence_ids

    # Verify generate_coverage_manifest
    main_slides = [{"finding_id": "EVID-EXEC-01", "order": 1, "title": "Exec Summary"}]
    app_slides = [{"finding_id": "EVID-GOV-01", "order": 8, "title": "Governance"}]
    manifest = generate_coverage_manifest(pkg["candidate_findings"], main_slides, app_slides)

    assert manifest["total_candidate_findings"] == len(pkg["candidate_findings"])
    assert manifest["main_deck_count"] == 1
    assert manifest["appendix_count"] == 1
    assert manifest["excluded_count"] == len(pkg["candidate_findings"]) - 2
    assert 0 <= manifest["coverage_pct"] <= 100
    for item in manifest["items"]:
        assert item["disposition"] in ("main_deck", "appendix", "excluded")
        assert len(item["reason"]) > 0


def test_voice_narration_readiness_model():
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)
    scope = {
        "scope_type": "workspace",
        "objective": "Voice Narration & Structural Readiness",
        "audience": "Executive Committee"
    }
    with get_connection() as conn:
        ws_ev = collect_workspace_evidence(conn, scope)
    deck_spec = generate_presentation_deck_spec(scope, ws_ev["primary_ctx"], workspace_evidence=ws_ev)

    assert len(deck_spec["slides"]) >= 8
    total_slides = len(deck_spec["slides"])

    for idx, slide in enumerate(deck_spec["slides"]):
        # Verify stable_slide_id
        assert "stable_slide_id" in slide
        assert isinstance(slide["stable_slide_id"], str) and len(slide["stable_slide_id"]) > 0

        # Verify narration_script
        assert "narration_script" in slide
        script = slide["narration_script"]
        assert isinstance(script, str) and len(script) > 20
        assert "**" not in script  # Free of markdown syntax

        # Verify reading_order
        assert "reading_order" in slide
        assert isinstance(slide["reading_order"], list)
        assert len(slide["reading_order"]) >= 4

        # Verify timing_metadata
        assert "timing_metadata" in slide
        tm = slide["timing_metadata"]
        assert "word_count" in tm and tm["word_count"] > 0
        assert "estimated_seconds" in tm and tm["estimated_seconds"] > 0

        # Verify order and total_slides
        assert slide["order"] == idx + 1
        assert slide["total_slides"] == total_slides


def test_new_presentation_layouts_pptx_export():
    test_deck = {
        "id": "deck_new_layouts_test",
        "theme": {"id": "corporate_navy"},
        "slides": [
            {
                "id": "s_action",
                "order": 1,
                "total_slides": 3,
                "layout": "action_plan",
                "title": "Strategic Operational Initiatives",
                "category": "STRATEGIC PROPOSALS",
                "subtitle": "Execution Roadmaps & Owners",
                "narrative": "Three immediate interventions targeted to address supply elasticity.",
                "structured_proposals": [
                    {
                        "priority": "HIGH",
                        "proposed_response": "Shift Schedule Optimization",
                        "motivating_finding": "High absenteeism during peak seasonal cycles",
                        "owner_role": "Unassigned - Operational Lead",
                        "success_metric": "Reduce absenteeism to < 3.5%",
                        "dependencies": "HRIS Automation"
                    },
                    {
                        "priority": "MEDIUM",
                        "proposed_response": "Cross-Store Inventory Rebalancing",
                        "motivating_finding": "Significant variance in weekly store throughput",
                        "owner_role": "Unassigned - Supply Chain Lead",
                        "success_metric": "+12% Stock Availability",
                        "dependencies": "Warehouse Logistics"
                    }
                ],
                "speaker_notes": "Present initiatives and ask for board sponsors."
            },
            {
                "id": "s_dual",
                "order": 2,
                "total_slides": 3,
                "layout": "two_charts",
                "title": "Dual Perspective Comparative Review",
                "category": "PERFORMANCE DYNAMICS",
                "subtitle": "Weekly Sales vs Labor Allocation",
                "narrative": "Comparison between store volume and staffing level across weeks.",
                "chart_left": {
                    "chart_type": "column",
                    "title": "Weekly Sales ($K)",
                    "categories": ["Wk 1", "Wk 2", "Wk 3"],
                    "series": [{"name": "Sales", "values": [120.0, 150.0, 180.0]}]
                },
                "chart_right": {
                    "chart_type": "line",
                    "title": "Staff Headcount",
                    "categories": ["Wk 1", "Wk 2", "Wk 3"],
                    "series": [{"name": "FTE", "values": [25.0, 28.0, 32.0]}]
                },
                "speaker_notes": "Dual chart comparison highlighting correlation."
            },
            {
                "id": "s_hero",
                "order": 3,
                "total_slides": 3,
                "layout": "full_chart_takeaway",
                "title": "Enterprise Sales Distribution (Hero)",
                "category": "MACRO TRAJECTORY",
                "subtitle": "Widescreen Focus",
                "narrative": "Full canvas visual emphasizing long-term store distribution trend.",
                "chart": {
                    "chart_type": "bar",
                    "title": "Store Performance",
                    "categories": ["North", "South", "East", "West"],
                    "series": [{"name": "Total Revenue", "values": [450.0, 320.0, 280.0, 190.0]}]
                },
                "speaker_notes": "Hero chart takeaway notes."
            }
        ]
    }

    pptx_path = export_spec_to_pptx(test_deck)
    assert pptx_path.exists()
    assert pptx_path.stat().st_size > 10000

    prs = Presentation(str(pptx_path))
    assert len(prs.slides) == 3

    # Slide 1: action_plan - check shapes and text
    s1 = prs.slides[0]
    s1_text = " ".join([shape.text_frame.text for shape in s1.shapes if shape.has_text_frame])
    assert "Shift Schedule Optimization" in s1_text
    assert "Unassigned - Operational Lead" in s1_text
    assert "HIGH PRIORITY" in s1_text
    assert "Slide 1 of 3" in s1_text

    # Slide 2: two_charts - check 2 native chart shapes
    s2 = prs.slides[1]
    chart_shapes_s2 = [shape for shape in s2.shapes if shape.has_chart]
    assert len(chart_shapes_s2) == 2
    s2_text = " ".join([shape.text_frame.text for shape in s2.shapes if shape.has_text_frame])
    assert "Slide 2 of 3" in s2_text

    # Slide 3: full_chart_takeaway - check 1 native chart shape
    s3 = prs.slides[2]
    chart_shapes_s3 = [shape for shape in s3.shapes if shape.has_chart]
    assert len(chart_shapes_s3) == 1
    s3_text = " ".join([shape.text_frame.text for shape in s3.shapes if shape.has_text_frame])
    assert "Slide 3 of 3" in s3_text


def test_quality_auditor_table_split_and_bounded_repair():
    deck_spec = {
        "id": "deck_audit_test",
        "theme": {"id": "clean_light"},
        "slides": [
            {
                "id": "s_tbl",
                "order": 1,
                "layout": "table_detail",
                "title": "Governance Log",
                "subtitle": "Detailed audit trail",
                "narrative": "Comprehensive row data.",
                "table": {
                    "headers": ["ID", "Finding", "Status"],
                    "rows": [[f"REC-{i}", f"Finding Description {i}", "Active"] for i in range(1, 13)]
                }
            },
            {
                "id": "s_chart_overflow",
                "order": 2,
                "layout": "chart_narrative",
                "title": "Department Comparisons",
                "subtitle": "Chart with overly long label",
                "narrative": "Reviewing long category name.",
                "chart": {
                    "chart_type": "column",
                    "title": "Metrics",
                    "categories": ["VeryLongDepartmentNameExceeding18Chars", "Engineering", "Marketing"],
                    "series": [{"name": "Headcount", "values": [50.0, 120.0, 85.0]}]
                }
            }
        ]
    }

    audit_res = PresentationQualityAuditor.audit_deck_spec(deck_spec, evidence_ledger=[])
    assert audit_res["status"] == "repairable"
    issue_types = [i["type"] for i in audit_res["issues"]]
    assert "table_row_overflow" in issue_types
    assert "long_category_label" in issue_types
    assert "missing_narration_script" in issue_types

    repaired_deck = PresentationQualityAuditor.execute_bounded_repair(deck_spec, audit_res)

    # 12-row table split into Part 1 (7 rows) and Part 2 (5 rows) -> 3 total slides
    assert len(repaired_deck["slides"]) == 3
    s1 = repaired_deck["slides"][0]
    s2 = repaired_deck["slides"][1]
    s3 = repaired_deck["slides"][2]

    assert "Part 1" in s1["title"]
    assert len(s1["table"]["rows"]) <= 7
    assert "Part 2" in s2["title"]
    assert len(s2["table"]["rows"]) == 5

    # Long category label truncated
    cats = s3["chart"]["categories"]
    assert any("..." in c for c in cats)
    assert not any(len(c) > 18 for c in cats)

    # Voice narration scripts populated
    for s in repaired_deck["slides"]:
        assert len(s.get("narration_script", "")) > 10
        assert "timing_metadata" in s

    # Re-audit should now have zero critical issues
    re_audit = PresentationQualityAuditor.audit_deck_spec(repaired_deck, evidence_ledger=[])
    assert re_audit["critical_count"] == 0


def test_eight_stage_pipeline_execution_and_stage_transitions():
    client = TestClient(app)
    sheet_id = seed_test_sales_sheet(client)

    from app.services.presentation_service import job_manager as default_job_mgr
    scope = {
        "scope_type": "workspace",
        "objective": "Pipeline 8-Stage Lifecycle Verification",
        "theme_id": "corporate_navy"
    }
    job_id = default_job_mgr.create_job(scope)

    # Execute synchronously in test thread
    execute_presentation_pipeline_async(job_id, scope, manager=default_job_mgr)

    res = client.get(f"/api/presentations/jobs/{job_id}")
    assert res.status_code == 200
    job = res.json()
    assert job["status"] == "ready"
    assert job["stage"] == "ready"
    assert job["progress_pct"] == 100
    assert job["deck_id"] is not None

    deck = job["deck"]
    assert deck["id"] == job["deck_id"]
    assert len(deck["slides"]) >= 8
    assert "coverage_manifest" in deck
    assert "quality_audit" in deck
    assert deck["quality_audit"]["critical_count"] == 0


