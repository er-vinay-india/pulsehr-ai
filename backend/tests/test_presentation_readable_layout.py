import copy
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from app.services.presentation.slide_layout import resolve_slide, presenter_notes, SlideLayoutError
from app.services.report_generator import export_spec_to_pptx, export_spec_to_pdf


def deck(slide, theme='executive_dark'):
    return {'id': 'layout-regression', 'theme': {'id': theme}, 'slides': [slide]}


def all_text(slide):
    return ' '.join([s.text for s in slide.shapes if s.has_text_frame] +
                    [c.text for s in slide.shapes if s.has_table for row in s.table.rows for c in row.cells])


@pytest.mark.parametrize('theme', ['executive_dark', 'clean_light'])
def test_comparison_table_pages_preserve_every_row_and_column(theme):
    rows = [[f'Employee {i:03}', 'Review recorded working days before confirming an exception', str(i)] for i in range(19)]
    slide = {'title': 'Recorded exceptions', 'layout': 'comparison_split',
             'narrative': 'Review each recorded exception.',
             'table': {'headers': ['Employee', 'Finding', 'Days'], 'rows': rows}}
    original = copy.deepcopy(slide)
    plan = resolve_slide(slide)
    assert len(plan['pages']) > 1
    flattened = [r for page in plan['pages'] for b in page if b['kind'] == 'table' for r in b['rows']]
    assert flattened == rows
    prs = Presentation(export_spec_to_pptx(deck(slide, theme)))
    actual = [list(c.text for c in row.cells) for s in prs.slides for shape in s.shapes if shape.has_table for row in list(shape.table.rows)[1:]]
    assert actual == rows
    assert len(prs.slides) == len(plan['pages'])
    assert slide == original
    pdf = export_spec_to_pdf(deck(slide, theme))
    assert pdf.stat().st_size > 1000


def test_kpi_third_takeaway_and_empty_hero_use_full_width():
    slide = {'title': 'Recorded results', 'layout': 'kpi_summary', 'metrics': [{'label': 'Employees', 'value': '259'}],
             'bullets': ['Confirm the period.', 'Review missing values.', 'Confirm the working-day calendar.']}
    prs = Presentation(export_spec_to_pptx(deck(slide)))
    assert all(b in all_text(prs.slides[0]) for b in slide['bullets'])
    plan = resolve_slide({'title': 'Attendance review', 'layout': 'title_hero', 'narrative': 'Review office days and approved leave.'})
    body = next(b for b in plan['blocks'] if b.get('field') == 'narrative')
    assert body['width'] == 864
    assert body['size'] >= 18


def test_chart_uses_main_canvas_and_keeps_supporting_detail_in_notes():
    slide = {'title': 'Recorded office days', 'layout': 'chart_narrative', 'narrative': 'Compare the recorded weekly averages.',
             'bullets': ['Only recorded values enter the average.'], 'metrics': [{'label': 'Employees', 'value': '3'}],
             'speaker_notes': 'Check the denominator.',
             'chart': {'title': 'Office days', 'unit': 'days', 'categories': ['Week 1', 'Week 2'],
                       'series': [{'name': 'Days', 'values': [1.25, 3.5]}]}}
    plan = resolve_slide(slide)
    block = next(b for b in plan['blocks'] if b['kind'] == 'chart')
    assert block['width'] == 864 and block['height'] >= 300
    prs = Presentation(export_spec_to_pptx(deck(slide)))
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [1.25, 3.5]
    assert chart.plots[0].data_labels.font.size.pt == 14
    assert not chart.has_legend
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert 'Only recorded values' in notes and 'Employees: 3' in notes and 'Check the denominator.' in notes
    assert presenter_notes(slide, plan) in notes


def test_action_dependencies_export_without_invented_targets():
    slide = {'title': 'Follow-up actions', 'layout': 'action_plan', 'initiatives': [
        {'title': 'Confirm the calendar', 'finding': 'Working-day definitions need confirmation.', 'dependency': 'HR approval'}]}
    prs = Presentation(export_spec_to_pptx(deck(slide)))
    text = all_text(prs.slides[0])
    assert 'HR approval' in text and 'Owner to confirm' in text
    assert '15%' not in text and 'SLA' not in text
    assert 'Working-day definitions' in prs.slides[0].notes_slide.notes_text_frame.text


def test_preview_uses_exact_export_geometry_and_reports_overflow():
    client = TestClient(app)
    slide = {'title': 'Recorded results', 'layout': 'title_hero', 'narrative': 'Review the recorded figures.'}
    response = client.post('/api/presentations/layout-preview', json=slide)
    assert response.status_code == 200
    assert response.json()['plan'] == resolve_slide(slide)
    slide['narrative'] = 'Long explanation. ' * 500
    assert client.post('/api/presentations/layout-preview', json=slide).status_code == 422


def test_legacy_cached_pass_does_not_approve_wrong_percentage_or_majority():
    slide = {'title': 'Colour composition', 'layout': 'chart_narrative', 'narrative': 'White has a majority.',
             'chart': {'chart_type': 'donut', 'unit': '%', 'categories': ['White', 'Blue', 'Other'],
                       'series': [{'name': 'Count', 'values': [407, 321, 272]}]}}
    spec = deck(slide)
    spec['metadata'] = {'validation_summary': {'status': 'PASSED', 'discrepancies_flagged': 0}}
    for export in (export_spec_to_pptx, export_spec_to_pdf):
        with pytest.raises(ValueError, match='evidence refresh'):
            export(spec)


def test_continuation_heading_requires_all_parts():
    from app.services.presentation.content_validation import semantic_issues
    incomplete = {'slides': [{'title': 'Recorded findings (1/2)'}]}
    assert any('Continuation title' in v for v in semantic_issues(incomplete))
    incomplete['slides'].append({'title': 'Recorded findings (2/2)'})
    assert not semantic_issues(incomplete)


def test_actions_preserve_timing_status_and_list_dependencies():
    slide = {'title': 'Proposed follow-up', 'layout': 'action_plan', 'initiatives': [
        {'title': 'Confirm period', 'status': 'Proposed', 'timeline': 'After HR confirmation',
         'dependencies': ['Calendar', 'Policy'], 'rationale': 'Comparable periods are required.'}]}
    plan = resolve_slide(slide)
    row = next(b for b in plan['blocks'] if b['kind'] == 'table')['rows'][0]
    assert 'Proposed' in row[1] and 'After HR confirmation' in row[1]
    assert 'Calendar; Policy' in row[2]
    assert 'Comparable periods are required.' in presenter_notes(slide, plan)


@pytest.mark.parametrize('kind', ['pie', 'donut', 'line', 'bar', 'column'])
def test_pdf_chart_families_export_recorded_values(kind, monkeypatch):
    from reportlab import rl_config
    monkeypatch.setattr(rl_config, 'pageCompression', 0)
    slide = {'title': 'Recorded category values', 'layout': 'chart_narrative',
             'chart': {'type': kind, 'unit': 'rows', 'categories': ['First', 'Second', 'Third'],
                       'series': [{'name': 'Records', 'values': [407, 321, 272]}]}}
    path = export_spec_to_pdf(deck(slide))
    text = path.read_bytes().decode('latin1')
    assert all(value in text for value in ['First', 'Second', 'Third', '407', '321', '272'])
    if kind in ('pie', 'donut'):
        assert '40.7%' in text


def test_mixed_pdf_keeps_readable_table_pagination(monkeypatch):
    from reportlab import rl_config
    monkeypatch.setattr(rl_config, 'pageCompression', 0)
    rows = [[f'Record {i:03}', 'Review the recorded value before confirming follow-up'] for i in range(25)]
    spec = {'id': 'mixed-layout', 'slides': [
        {'title': 'Specialized introduction', 'layout': 'custom_existing_layout', 'narrative': 'Introduction'},
        {'title': 'Recorded findings', 'layout': 'comparison_split', 'table': {'headers': ['Record', 'Finding'], 'rows': rows}}]}
    text = export_spec_to_pdf(spec).read_bytes().decode('latin1')
    assert all(row[0] in text for row in rows)
    assert 'Slide 2 of 2' in text


def test_editor_pdf_endpoint_exports_current_verified_spec():
    from app.services.presentation.hr_report import build_hr_report
    from test_presentation_hr_report import fixture_context
    spec = build_hr_report({}, fixture_context())
    client = TestClient(app)
    response = client.post('/api/presentations/export-pdf', json={'deck_spec': spec})
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/pdf'
    assert response.content.startswith(b'%PDF')
    spec['slides'][0]['metrics'][0]['value'] = '99999'
    response = client.post('/api/presentations/export-pdf', json={'deck_spec': spec})
    assert response.status_code == 422


def test_native_composition_has_editable_chart_and_readable_value_key():
    slide = {'title': 'Recorded distribution', 'layout': 'chart_narrative',
             'chart': {'type': 'donut', 'unit': 'rows', 'categories': ['White', 'Blue', 'Other'],
                       'series': [{'name': 'Records', 'values': [407, 321, 272]}]}}
    prs = Presentation(export_spec_to_pptx(deck(slide)))
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [407, 321, 272]
    assert 'White: 407 (40.7%)' in all_text(prs.slides[0])
    assert not chart.plots[0].has_data_labels
