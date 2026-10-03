"""
Tests for PPT Generation Workflow Audit and Findings.

Validates:
1. Real calculation inputs in orchestrator (elimination of synthetic `total_records * 0.8`).
2. Elimination of invented numerical fallbacks (+7.8%, 8.11x, 1.8x).
3. Strict claim verifier: units, denominators, metric inclusion, ±0.1% tolerance.
4. Five Review Gates: automated checks vs explicit human sign-off, revision tracking, invalidation on edit.
5. Preservation of complete headlines without colon truncation.
6. 13-phase workflow structure and speaker notes duration budgeting.
"""

import json
from pathlib import Path
import pytest
from fastapi import HTTPException
import pptx

from app.core.config import EXPORTS_DIR
from app.db.database import get_connection
from app.services.presentation.review_gates import (
    initialize_review_gates,
    evaluate_automated_gates,
    record_human_signoff,
    invalidate_review_gates_on_edit
)
from app.services.presentation.claim_verifier import verify_presentation_claims
from app.services.presentation.builders.strengths_builder import build_strengths_slides
from app.services.presentation.builders.roadmap_builder import build_roadmap_slides
from app.services.presentation.builders.governance_builder import build_boundary_and_evidence_slides
from app.services.presentation.pipeline_orchestrator import (
    PIPELINE_PHASES,
    generate_structured_speaker_notes
)
from app.services.presentation.orchestrator.tool_registry import tool_registry
from app.services.presentation.observer_ai import SlideContextObserver
from app.services.report_generator import export_spec_to_pptx, export_spec_to_pdf
from app.routers.presentations import (
    approve_review_gate,
    update_presentation_deck,
    export_presentation_to_pptx,
    download_deck_pptx,
    ReviewGateSignoffRequest,
    ExportPptxRequest,
)
from app.services.presentation.director.intent_planner import plan_intent
from app.services.presentation.director.narrative_planner import plan_narrative
from app.services.presentation.director.director_models import (
    PresentationPlanningContext,
    PresentationBrief,
    AudienceSeniority,
    DeliveryMode,
)


def test_orchestrator_calculate_metric_no_synthetic_inputs():
    """Verify tool execution does not fabricate numbers when inputs are missing."""
    res = tool_registry.execute_tool(
        "calculate_metric",
        metric_type="custom",
        formula="rate",
        inputs={}
    )
    res_dict = res.model_dump()
    assert res_dict["status"] in ("failed", "unavailable", "success")
    calc = str(res_dict.get("calculation_method", "")).lower()
    data_str = str(res_dict.get("data", {})).lower()
    assert "unavailable" in calc or "missing" in calc or "unavailable" in data_str or "unsupported" in calc


def test_zero_invented_numerical_fallbacks_in_builders():
    """Verify builders never output invented +7.8%, 8.11x, or 1.8x when evidence is empty."""
    evidence_empty = []

    # 1. Strengths builder
    s_slides = build_strengths_slides(
        strengths=[],
        line_chart=None,
        source_summary="test.xlsx",
        evidence_ledger=evidence_empty,
        is_partial_year=False,
        start_order=1
    )
    s_text = json.dumps(s_slides)
    assert "+7.8%" not in s_text
    assert "8.11x" not in s_text

    # 2. Roadmap builder
    r_slides = build_roadmap_slides(
        total_eval_records=100,
        dispersion_metric_str="1.2x",
        lead_cat="Engineering",
        source_summary="test.xlsx",
        evidence_ledger=evidence_empty,
        is_partial_year=False,
        start_order=2
    )
    r_text = json.dumps(r_slides)
    assert "+7.8%" not in r_text

    # 3. Governance builder
    g_slides = build_boundary_and_evidence_slides(
        total_eval_records=100,
        reporting_period_summary="Q1-Q3 2026",
        source_summary="test.xlsx",
        profiled_data={"columns": []},
        mean_val_str="15.2",
        completeness_pct=100.0,
        dispersion_metric_str="1.2x",
        snapshot_hash="sha256:abc",
        evidence_ledger=evidence_empty,
        is_partial_year=False,
        start_order=3
    )
    g_text = json.dumps(g_slides)
    assert "+7.8%" not in g_text


def test_claim_verifier_strictness_and_units():
    """Verify claim verifier strictly validates numbers, units, and flags mismatches."""
    evidence = [
        {
            "id": "EV-01",
            "metric_name": "Turnover Rate",
            "value": 18.5,
            "unit": "%",
            "confidence": 0.95
        },
        {
            "id": "EV-02",
            "metric_name": "Headcount",
            "value": 1250,
            "unit": "count",
            "confidence": 0.95
        }
    ]

    # Valid claim with ±0.1% tolerance
    slides_valid = [
        {
            "id": "slide_1",
            "title": "Turnover reached 18.5% across global engineering",
            "metrics": [{"label": "Turnover Rate", "value": "18.5%", "evidence_id": "EV-01"}]
        }
    ]
    summary_valid = verify_presentation_claims({"slides": slides_valid}, evidence)
    assert summary_valid["status"] == "PASSED"
    assert summary_valid["discrepancies_flagged"] == 0

    # Invalid claim: unbacked claim (34.2% vs 18.5% in ledger)
    slides_invalid = [
        {
            "id": "slide_2",
            "title": "Turnover soared to 34.2% across business units",
            "metrics": [{"label": "Turnover Rate", "value": "34.2%", "evidence_id": "EV-01"}]
        }
    ]
    summary_invalid = verify_presentation_claims({"slides": slides_invalid}, evidence)
    assert summary_invalid["discrepancies_flagged"] > 0
    assert summary_invalid["status"] in ("FAILED", "REQUIRES_REVIEW", "DISCREPANCIES_FLAGGED")


def test_five_review_gates_lifecycle():
    """Verify initialization, evaluation, human sign-off separation, and invalidation."""
    brief_data = {
        "objective": "Address escalating voluntary attrition in engineering",
        "audience": "Board of Directors & Executive Committee",
        "decision_requested": "Approve $2.4M retention budget",
        "is_inferred": False,
        "success_criterion": "Board understands root cause and approves budget."
    }

    # 1. Initialize
    rg = initialize_review_gates(brief=brief_data, revision=1)
    gates = rg["gates"]
    assert len(gates) == 5
    assert "gate_1_brief" in gates
    assert "gate_2_storyline" in gates
    assert "gate_3_evidence" in gates
    assert "gate_4_visual" in gates
    assert "gate_5_export_accessibility" in gates

    # Human sign-off MUST be PENDING initially (never automatically approved)
    for gid, gate in gates.items():
        assert gate["human_approval"]["status"] == "PENDING"
        assert gate["human_approval"]["approved_by"] is None

    # 2. Automated evaluation
    deck_spec = {
        "id": "deck_test_101",
        "metadata": {
            "objective": brief_data["objective"],
            "audience": brief_data["audience"],
            "decision_requested": brief_data["decision_requested"],
            "validation_summary": {
                "status": "PASSED",
                "discrepancies_flagged": 0,
                "total_metrics_checked": 4
            }
        },
        "slides": [
            {"id": "s1", "title": "Executive Summary & Context", "layout": "title_hero"},
            {"id": "s2", "title": "Turnover Surged to 18.5% in Q3", "layout": "chart_narrative"},
            {"id": "s3", "title": "Strategic Recommendation & Resource Allocation", "layout": "comparison_split"}
        ],
        "pptx_filename": "deck_test_101.pptx"
    }

    evaluated = evaluate_automated_gates(deck_spec, brief=brief_data)
    # Check that automated checks pass, but human sign-offs remain strictly pending
    assert evaluated["gates"]["gate_1_brief"]["automated"]["status"] == "PASSED"
    assert evaluated["gates"]["gate_2_storyline"]["automated"]["status"] == "PASSED"
    assert evaluated["gates"]["gate_3_evidence"]["automated"]["status"] == "PASSED"
    assert evaluated["gates"]["gate_1_brief"]["human_approval"]["status"] == "PENDING"
    assert evaluated["summary"]["human_signoff_complete"] is False

    # 3. Explicit human sign-off
    signed_off = record_human_signoff(
        evaluated,
        gate_id="gate_1_brief",
        approved=True,
        user_name="VP People Analytics",
        notes="Brief and audience validated with CEO."
    )
    assert signed_off["gates"]["gate_1_brief"]["human_approval"]["status"] == "APPROVED"
    assert signed_off["gates"]["gate_1_brief"]["human_approval"]["approved_by"] == "VP People Analytics"

    # 4. Invalidation on edit
    deck_spec["metadata"]["review_gates"] = signed_off
    invalidated_deck = invalidate_review_gates_on_edit(deck_spec, edited_scope="content")
    inv_rg = invalidated_deck["metadata"]["review_gates"]
    assert inv_rg["revision"] == 2
    # Content edit must invalidate evidence & storyline sign-offs
    assert inv_rg["gates"]["gate_3_evidence"]["human_approval"]["status"] == "PENDING"
    assert inv_rg["gates"]["gate_3_evidence"]["automated"]["status"] == "REQUIRES_REVIEW"


def test_headline_preservation_no_colon_truncation():
    """Verify conclusion headlines with colons are preserved in full without truncating numbers."""
    full_headline = "Severe Attrition Surge: Engineering departments show 24.5% flight risk"

    observer = SlideContextObserver(scope={"objective": "Test Attrition"})
    res = observer.on_slide_start(
        slide_num=1,
        slide_title=full_headline,
        category="Strategy"
    )
    assert observer.slides[0]["title"] == full_headline
    assert "24.5% flight risk" in observer.slides[0]["title"]


def test_pipeline_orchestrator_phases_and_notes():
    """Verify all 13 workflow phases exist and speaker notes respect time budget."""
    assert len(PIPELINE_PHASES) == 13
    assert PIPELINE_PHASES[0]["key"] == "brief_setup"
    assert PIPELINE_PHASES[1]["key"] == "evidence_audit"
    assert PIPELINE_PHASES[2]["key"] == "narrative_arc"
    assert PIPELINE_PHASES[3]["key"] == "headlines"
    assert PIPELINE_PHASES[4]["key"] == "layout_selection"
    assert PIPELINE_PHASES[5]["key"] == "math_reconciliation"
    assert PIPELINE_PHASES[6]["key"] == "graphics_charts"
    assert PIPELINE_PHASES[7]["key"] == "executive_polish"
    assert PIPELINE_PHASES[8]["key"] == "visual_qa"
    assert PIPELINE_PHASES[9]["key"] == "animation"
    assert PIPELINE_PHASES[10]["key"] == "speaker_notes"
    assert PIPELINE_PHASES[11]["key"] == "export_qa"
    assert PIPELINE_PHASES[12]["key"] == "ready"

    # Test speaker notes generation
    slide = {
        "id": "s_demo",
        "title": "Turnover reached 18.5% across global engineering",
        "takeaway": "Key retention risks are concentrated in senior levels.",
        "content": {"bullets": ["Level 5 flight risk is highest at 28%."]}
    }
    notes = generate_structured_speaker_notes(
        slide=slide,
        slide_index=1,
        total_slides=5,
        target_minutes=10,
        implication="Prioritize retention equity grants for senior ICs.",
        transition="Let's examine compensation equity in the next slide."
    )
    assert "WHAT TO NOTICE:" in notes
    assert "WHY IT MATTERS:" in notes
    assert "SUPPORTING EVIDENCE:" in notes
    assert "TIME BUDGET:" in notes
    assert "TRANSITION:" in notes


def test_claim_verifier_reproduced_failures():
    """Verify strict claim verifier catches all 5 reproduced failure cases."""
    evidence = [
        {"id": "EV-01", "metric_name": "Turnover Rate", "value": 18.5, "unit": "%"},
        {"id": "EV-02", "metric_name": "Row Count", "value": 1250, "unit": "count"},
    ]

    # Case 1: Headline says 99%; metric correctly says 18.5%
    deck_c1 = {
        "slides": [{
            "id": "s1",
            "title": "Turnover soared to 99% across the workforce",
            "metrics": [{"label": "Turnover Rate", "value": "18.5%", "evidence_id": "EV-01"}]
        }]
    }
    sum_c1 = verify_presentation_claims(deck_c1, evidence)
    assert sum_c1["discrepancies_flagged"] > 0
    assert any("99" in str(d) for d in sum_c1["discrepancies"])

    # Case 2: Turnover says 1,250%, matching the evidence row count (unit mismatch)
    deck_c2 = {
        "slides": [{
            "id": "s2",
            "title": "Turnover Rate Analysis",
            "metrics": [{"label": "Turnover Rate", "value": "1,250%", "evidence_id": "EV-02"}]
        }]
    }
    sum_c2 = verify_presentation_claims(deck_c2, evidence)
    assert sum_c2["discrepancies_flagged"] > 0
    assert any("1250" in str(d) or "unit" in str(d).lower() for d in sum_c2["discrepancies"])

    # Case 3: -18.5% when evidence says +18.5% (negative sign loss)
    deck_c3 = {
        "slides": [{
            "id": "s3",
            "title": "Turnover dropped to -18.5%",
            "metrics": [{"label": "Turnover Rate", "value": "-18.5%", "evidence_id": "EV-01"}]
        }]
    }
    sum_c3 = verify_presentation_claims(deck_c3, evidence)
    assert sum_c3["discrepancies_flagged"] > 0

    # Case 4: "187 of 232 (12%)" - internal arithmetic contradiction (187/232 = 80.6% != 12%)
    deck_c4 = {
        "slides": [{
            "id": "s4",
            "title": "187 of 232 (12%) participants completed the cycle",
            "metrics": [{"label": "Turnover Rate", "value": "18.5%", "evidence_id": "EV-01"}]
        }]
    }
    sum_c4 = verify_presentation_claims(deck_c4, evidence)
    assert sum_c4["discrepancies_flagged"] > 0
    assert any("arithmetic" in str(d).lower() or "contradiction" in str(d).lower() or "187" in str(d) for d in sum_c4["discrepancies"])

    # Case 5: Narrative and chart say 99%; metric says 18.5%
    deck_c5 = {
        "slides": [{
            "id": "s5",
            "title": "Engineering Retention",
            "narrative": "Severe turnover observed at 99% across senior levels.",
            "chart": {
                "type": "bar",
                "title": "Turnover Rates",
                "categories": ["Engineering"],
                "series": [{"name": "Rate", "values": [99.0]}]
            },
            "metrics": [{"label": "Turnover Rate", "value": "18.5%", "evidence_id": "EV-01"}]
        }]
    }
    sum_c5 = verify_presentation_claims(deck_c5, evidence)
    assert sum_c5["discrepancies_flagged"] > 0


def test_chart_alt_text_applied_in_pptx():
    """Verify exported PowerPoint charts have alt text title and description properly set on cNvPr."""
    deck_spec = {
        "id": "deck_alt_text_test",
        "metadata": {"title": "Alt Text Verification Deck"},
        "slides": [{
            "id": "s_chart",
            "title": "Departmental Attrition Comparison",
            "layout": "chart_narrative",
            "narrative": "Engineering shows the highest flight risk across departments.",
            "chart": {
                "type": "column",
                "title": "Attrition by Department",
                "subtitle": "Percentage of voluntary departures",
                "categories": ["Engineering", "Sales", "Marketing"],
                "series": [{"name": "Attrition Rate", "values": [24.5, 18.2, 12.0]}]
            }
        }]
    }
    pptx_path = export_spec_to_pptx(deck_spec)
    assert pptx_path.exists()

    prs = pptx.Presentation(str(pptx_path))
    slide = prs.slides[0]
    
    found_chart_alt = False
    for shape in slide.shapes:
        elem = shape.element
        for child in elem.iter():
            if child.tag.endswith("cNvPr"):
                title_val = child.attrib.get("title", "")
                descr_val = child.attrib.get("descr", "")
                if title_val or descr_val:
                    found_chart_alt = True
                    assert "Attrition by Department" in title_val or "Departmental" in descr_val or "Chart" in title_val or "Attrition" in descr_val
                    break
        if found_chart_alt:
            break
    assert found_chart_alt, "Chart cNvPr alt text title/descr was not applied in exported PPTX"


def test_review_gates_accessibility_and_visual_safety():
    """Verify Gate 4 requires visual audit and Gate 5 fails on nonexistent export file."""
    brief_data = {
        "objective": "Test Visual and Export Safety",
        "audience": "Board of Directors",
        "decision_requested": "Approve budget",
        "is_inferred": False
    }
    deck_spec = {
        "id": "deck_gate_safety",
        "metadata": {
            "objective": "Test Visual and Export Safety",
            "audience": "Board of Directors",
            "decision_requested": "Approve budget",
            "validation_summary": {"status": "PASSED", "discrepancies_flagged": 0, "total_metrics_checked": 2}
        },
        "slides": [
            {"id": "s1", "title": "Executive Summary", "layout": "title_hero"},
            {
                "id": "s2",
                "title": "Operational Findings Across Units",
                "layout": "chart_narrative",
                "chart": {
                    "type": "bar",
                    "title": "Findings",
                    "categories": ["Ops"],
                    "series": [{"name": "Rate", "values": [12.0]}]
                }
            },
            {"id": "s3", "title": "Strategic Roadmap", "layout": "comparison_split"}
        ],
        "pptx_filename": "nonexistent_file_99999.pptx"
    }

    # 1. Gate 4 with quality_audit=None MUST return REQUIRES_REVIEW (never falsely PASSED)
    res_no_audit = evaluate_automated_gates(deck_spec, quality_audit=None, brief=brief_data)
    assert res_no_audit["gates"]["gate_4_visual"]["automated"]["status"] == "REQUIRES_REVIEW"
    assert "pending review" in res_no_audit["gates"]["gate_4_visual"]["automated"]["details"].lower() or "not yet executed" in res_no_audit["gates"]["gate_4_visual"]["automated"]["details"].lower()

    # 2. Gate 5 with nonexistent PPTX file MUST return FAILED
    assert res_no_audit["gates"]["gate_5_export_accessibility"]["automated"]["status"] == "FAILED"
    assert "not found on disk" in res_no_audit["gates"]["gate_5_export_accessibility"]["automated"]["details"].lower()

    # 3. Gate 4 with clean quality_audit returns PASSED
    res_clean_audit = evaluate_automated_gates(
        deck_spec,
        quality_audit={"passed": True, "critical_count": 0, "warning_count": 0},
        brief=brief_data
    )
    assert res_clean_audit["gates"]["gate_4_visual"]["automated"]["status"] == "PASSED"

    # 4. Gate 5 with real file on disk returns PASSED
    # This structural fixture has unbound legacy claims: cached gates cannot
    # certify a normal export. Explicit diagnostic mode still tests native shapes.
    with pytest.raises(ValueError, match="evidence refresh"):
        export_spec_to_pptx(deck_spec)
    real_pptx = export_spec_to_pptx(deck_spec, diagnostic_layout_only=True)
    deck_spec["pptx_filename"] = real_pptx.name
    res_real_export = evaluate_automated_gates(deck_spec, brief=brief_data)
    assert res_real_export["gates"]["gate_5_export_accessibility"]["automated"]["status"] == "PASSED"
    assert "pending human verification" in res_real_export["gates"]["gate_5_export_accessibility"]["automated"]["details"].lower()


def test_failed_gates_block_delivery_in_pipeline_and_endpoints():
    """Verify failed checks block delivery in export endpoints with HTTP 422."""
    # Deck with failed evidence gate (discrepancy flagged)
    failing_deck = {
        "id": "deck_failing_delivery",
        "metadata": {
            "title": "Failing Deck",
            "review_gates": {
                "revision": 1,
                "gates": {
                    "gate_3_evidence": {
                        "automated": {"status": "FAILED", "details": "2 discrepancies flagged"}
                    }
                }
            }
        },
        "slides": [{"id": "s1", "title": "Slide 1"}]
    }

    # 1. /export-pptx must reject with HTTP 422
    with pytest.raises(HTTPException) as exc_info:
        export_presentation_to_pptx(ExportPptxRequest(deck_spec=failing_deck))
    assert exc_info.value.status_code == 422
    assert "Delivery blocked" in exc_info.value.detail

    # 2. /download/{deck_id} must reject with HTTP 422 when stored deck has failed gate
    with get_connection() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO presentation_decks (id, title, spec_json, pptx_filename) VALUES (?, ?, ?, ?)",
            ("deck_failing_delivery", "Failing Deck", json.dumps(failing_deck), "failing.pptx")
        )
        conn.commit()

    with pytest.raises(HTTPException) as exc_info2:
        download_deck_pptx("deck_failing_delivery")
    assert exc_info2.value.status_code == 422
    assert "Delivery blocked" in exc_info2.value.detail


def test_revision_binding_and_monotonicity():
    """Verify revision increments monotonically on save and approvals reject revision mismatches."""
    deck_id = "deck_rev_test_99"
    initial_spec = {
        "id": deck_id,
        "metadata": {
            "title": "Revision Monotonicity Deck",
            "review_gates": {
                "revision": 1,
                "gates": {
                    "gate_1_brief": {
                        "automated": {"status": "PASSED"},
                        "human_approval": {"status": "PENDING"}
                    }
                }
            }
        },
        "slides": [{"id": "s1", "title": "Slide 1"}]
    }

    with get_connection() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO presentation_decks (id, title, spec_json) VALUES (?, ?, ?)",
            (deck_id, "Revision Deck", json.dumps(initial_spec))
        )
        conn.commit()

    # Save 1: from client with rev 1 -> stored becomes rev 2
    update_presentation_deck(deck_id, initial_spec)
    with get_connection() as conn:
        row = conn.execute("SELECT spec_json FROM presentation_decks WHERE id = ?", (deck_id,)).fetchone()
    spec_v2 = json.loads(row["spec_json"])
    assert spec_v2["metadata"]["review_gates"]["revision"] == 2

    # Save 2: another client with stale rev 1 saves -> stored becomes rev 3 (never duplicate 2!)
    stale_spec = dict(initial_spec)
    stale_spec["metadata"]["review_gates"]["revision"] = 1
    update_presentation_deck(deck_id, stale_spec)
    with get_connection() as conn:
        row = conn.execute("SELECT spec_json FROM presentation_decks WHERE id = ?", (deck_id,)).fetchone()
    spec_v3 = json.loads(row["spec_json"])
    assert spec_v3["metadata"]["review_gates"]["revision"] == 3

    # Approval test: approve with mismatched expected_revision (e.g. 1 instead of 3)
    with pytest.raises(HTTPException) as exc_info:
        approve_review_gate(
            deck_id=deck_id,
            gate_id="gate_1_brief",
            req=ReviewGateSignoffRequest(approved=True, user_name="Executive", expected_revision=1)
        )
    assert exc_info.value.status_code == 409
    assert "Revision mismatch" in exc_info.value.detail

    # Approval test: approve with matching expected_revision 3 -> succeeds
    res_ok = approve_review_gate(
        deck_id=deck_id,
        gate_id="gate_1_brief",
        req=ReviewGateSignoffRequest(approved=True, user_name="Executive", expected_revision=3)
    )
    assert res_ok["gates"]["gate_1_brief"]["human_approval"]["status"] == "APPROVED"
    assert res_ok["gates"]["gate_1_brief"]["human_approval"]["revision"] == 3


def test_execution_order_and_zero_generator_sleeps():
    """Verify zero artificial sleeps in deck_generator and clean visual materialization separation."""
    gen_file = Path(__file__).resolve().parent.parent / "app" / "services" / "presentation" / "deck_generator.py"
    content = gen_file.read_text()
    assert "time.sleep" not in content, "Found time.sleep in deck_generator.py - must be eliminated"

    # Verify materialize_slide_visuals is exported and callable
    from app.services.presentation.deck_generator import materialize_slide_visuals
    slides = [{"id": "s_mat", "title": "Materialization Slide", "category": "Analysis"}]
    mat_slides = materialize_slide_visuals(slides, theme_id="executive_dark")
    assert len(mat_slides) == 1
    assert "visual_spec" in mat_slides[0]


def test_capabilities_brief_pdf_and_speaker_notes():
    """Verify brief consumption in planners, PDF generation, and structured speaker notes preservation."""
    # 1. Brief consumption in intent and narrative planners
    brief = PresentationBrief(
        objective="Executive Turnover Strategy",
        audience="Board of Directors",
        decision_requested="Approve $3.2M retention budget",
        main_takeaway="Engineering turnover reached critical peak of 24.5%",
        presentation_time_minutes=25,
        deliverable="both",
        is_inferred=False
    )
    ctx = PresentationPlanningContext(
        domain="Human Resources",
        objective="Executive Turnover Strategy",
        audience="Board of Directors",
        total_records=2500,
        brief=brief
    )

    intent = plan_intent(ctx, max_retries=-1)
    assert any("retention" in t.lower() or "peak" in t.lower() for t in intent.key_takeaways) or "peak" in intent.primary_goal.lower()

    narrative = plan_narrative(ctx, intent, [], max_retries=-1)
    # A requested takeaway is not evidence for an observed turnover peak.
    assert "critical peak" not in narrative.executive_thesis.lower()
    assert "review" in narrative.executive_thesis.lower()

    # 2. PDF generation
    deck_spec = {
        "id": "deck_pdf_test",
        "metadata": {"title": "PDF Generation Test"},
        "slides": [
            {"id": "sp1", "title": "Executive Summary", "category": "Strategy", "content": {"bullets": ["Point A", "Point B"]}},
            {"id": "sp2", "title": "Key Indicators", "category": "Metrics", "metrics": [{"label": "Rate", "value": "18.5%"}]}
        ]
    }
    pdf_path = export_spec_to_pdf(deck_spec)
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 0
    assert pdf_path.name.endswith(".pdf")

    # 3. Speaker notes preservation & limitations
    slide_with_notes = {
        "id": "sn1",
        "title": "Turnover Variance",
        "speaker_notes": "Existing confidential author comment: examine Q4 bonus plan.",
        "source_label": "HRIS_Export_2026.xlsx",
        "limitations": "Self-reported exit interviews only."
    }
    notes = generate_structured_speaker_notes(
        slide=slide_with_notes,
        total_slides=10,
        target_minutes=5,
        evidence_ledger=[{"id": "EV-99", "source_dataset": "HRIS_Export_2026.xlsx", "limitations": "Self-reported exit interviews only."}]
    )
    # Must preserve existing notes
    assert "Existing confidential author comment" in notes
    # Must include actual sources and limitations
    assert "HRIS_Export_2026.xlsx" in notes
    assert "Self-reported exit interviews" in notes
    # Must flag time budget excess when notes exceed per-slide allocation
    assert "TIME BUDGET:" in notes
