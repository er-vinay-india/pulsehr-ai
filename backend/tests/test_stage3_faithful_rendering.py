"""Stage 3 Focused Unit Tests: Faithful PPTX Rendering & Visual Reuse Invariants."""

import pytest
import pptx
from pptx.util import Inches

from app.services.presentation.content_models import (
    ContentKind,
    ContentRole,
    PresentationContentItem,
    VisualStrategy,
)
from app.services.presentation.story_models import ContentBindingRole
from app.services.presentation.layer_zone_models import (
    CoordinateSpace,
    EditabilityPolicy,
    FrameInheritancePolicy,
    LayeredSlideSpec,
    SlideObjectSpec,
    StructureStatus,
    ZoneSpec,
    ZoneType,
)
from app.services.presentation.layout_models import (
    Box2D,
    ResolvedObjectLayout,
    ResolvedSafeArea,
)
from app.services.presentation.render_models import (
    RenderMode,
    ShapeType,
)
from app.services.presentation.renderer_adapters import (
    ChartRendererAdapter,
    RasterRendererAdapter,
    TableRendererAdapter,
)
from app.services.presentation.renderer_validator import (
    RendererValidationError,
    RendererValidator,
)


def _make_dummy_slide():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return prs, slide


def _make_dummy_slide_spec(slide_id="SLIDE_TEST_01"):
    return LayeredSlideSpec(
        physical_slide_id=slide_id,
        logical_slide_id="LOG_01",
        semantic_topic="TEST_TOPIC",
        structural_role="DIAGNOSTIC_ANALYSIS",
        sequence_index=1,
        continuation_index=1,
        is_continuation=False,
        headline="Test Headline",
        zones=[],
        structure_status=StructureStatus.STRUCTURE_VALIDATED,
    )


def test_choropleth_renders_as_preserved_image():
    """Choropleth must render as PRESERVED_IMAGE with shape_type PICTURE."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-MAP-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Geographic Density",
        visual_id="TOPIC-MAP-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="choropleth",
        source_surface="dashboard",
        kind=ContentKind.MAP,
        title="India Population Density Choropleth",
        narrative_takeaway="Delhi commands highest density",
        visual_spec={
            "chart_type": "choropleth",
            "categories": ["Delhi", "Maharashtra", "Arunachal Pradesh"],
            "values": [11320.0, 365.0, 17.0],
            "unit": "people/sq. km",
            "geo_scope": "india",
        },
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_MAP_01",
        object_name="map_object",
        content_id=item.content_id,
        name="Map Object",
        kind=ContentKind.MAP,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.RASTER_ALLOWED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = RasterRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert record.render_mode == RenderMode.PRESERVED_IMAGE
    assert record.rendered_visual_family == "choropleth"
    assert record.shape_type == ShapeType.PICTURE
    assert record.is_native_editable is False
    assert record.source_visual_id == "TOPIC-MAP-01"


def test_scatter_renders_as_native_scatter():
    """Scatter must render as NATIVE PowerPoint XY Scatter chart."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-SCAT-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Correlation",
        visual_id="TOPIC-SCAT-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="scatter",
        source_surface="dashboard",
        kind=ContentKind.CHART,
        title="Population vs Decadal Growth Rate",
        narrative_takeaway="r = -0.04 (Negligible association)",
        visual_spec={
            "chart_type": "scatter",
            "scatter_points": [
                {"x": 199812341.0, "y": 20.2, "label": "Uttar Pradesh"},
                {"x": 112374333.0, "y": 16.0, "label": "Maharashtra"},
                {"x": 1383727.0, "y": 26.0, "label": "Arunachal Pradesh"},
            ],
            "x_label": "Population 2011",
            "y_label": "Growth Rate",
            "unit": "%",
        },
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_SCAT_01",
        object_name="scatter_object",
        content_id=item.content_id,
        name="Scatter Object",
        kind=ContentKind.CHART,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.NATIVE_REQUIRED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = ChartRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert getattr(shape, "has_chart", False) is True
    assert record.render_mode == RenderMode.NATIVE
    assert record.rendered_visual_family == "scatter"
    assert record.shape_type == ShapeType.CHART
    assert record.is_native_editable is True
    assert record.source_visual_id == "TOPIC-SCAT-01"


def test_box_plot_renders_as_preserved_image():
    """Box plot must render as PRESERVED_IMAGE because PowerPoint has no native DrawingML box plot."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-BOX-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Distribution",
        visual_id="TOPIC-BOX-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="box_plot",
        source_surface="dashboard",
        kind=ContentKind.CHART,
        title="Population Spread & Distribution",
        narrative_takeaway="Distribution spans interquartile range with 67.2% metric",
        visual_spec={
            "chart_type": "box_plot",
            "values": [20.0, 16.0, 14.0, 11.0, 8.0],
            "unit": "%",
            "key_metric": "67.2%",
        },
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_BOX_01",
        object_name="box_object",
        content_id=item.content_id,
        name="Box Object",
        kind=ContentKind.CHART,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.RASTER_ALLOWED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = ChartRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert record.render_mode == RenderMode.PRESERVED_IMAGE
    assert record.rendered_visual_family == "box_plot"
    assert record.shape_type == ShapeType.PICTURE
    assert record.is_native_editable is False
    assert record.source_visual_id == "TOPIC-BOX-01"


def test_variance_bar_renders_as_native_bar():
    """Variance bar must render as NATIVE PowerPoint clustered bar chart."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-BAR-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Anomalies",
        visual_id="TOPIC-BAR-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="variance_bar",
        source_surface="dashboard",
        kind=ContentKind.CHART,
        title="Population Outliers & Variances",
        narrative_takeaway="Top: Uttar Pradesh (199.8M units)",
        visual_spec={
            "chart_type": "variance_bar",
            "categories": ["Uttar Pradesh", "Maharashtra", "Bihar"],
            "values": [199812341.0, 112374333.0, 104099452.0],
            "benchmark": 33634860.47,
            "unit": "units",
        },
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_BAR_01",
        object_name="bar_object",
        content_id=item.content_id,
        name="Bar Object",
        kind=ContentKind.CHART,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.NATIVE_REQUIRED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = ChartRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert getattr(shape, "has_chart", False) is True
    assert record.render_mode == RenderMode.NATIVE
    assert record.rendered_visual_family == "variance_bar"
    assert record.shape_type == ShapeType.CHART
    assert record.is_native_editable is True
    assert record.source_visual_id == "TOPIC-BAR-01"


def test_bullet_renders_as_preserved_image():
    """Bullet benchmark visual must render as PRESERVED_IMAGE with target line."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-BUL-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Benchmarks",
        visual_id="TOPIC-BUL-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="bullet",
        source_surface="dashboard",
        kind=ContentKind.CHART,
        title="Entities Meeting Benchmark Target",
        narrative_takeaway="Target: 33,634,860 units",
        visual_spec={
            "chart_type": "bullet",
            "categories": ["Uttar Pradesh", "Maharashtra", "Bihar"],
            "values": [199812341.0, 112374333.0, 104099452.0],
            "benchmark": 33634860.47,
            "unit": "units",
        },
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_BUL_01",
        object_name="bullet_object",
        content_id=item.content_id,
        name="Bullet Object",
        kind=ContentKind.CHART,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.RASTER_ALLOWED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = ChartRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert record.render_mode == RenderMode.PRESERVED_IMAGE
    assert record.rendered_visual_family == "bullet"
    assert record.shape_type == ShapeType.PICTURE
    assert record.is_native_editable is False
    assert record.source_visual_id == "TOPIC-BUL-01"


def test_table_renders_as_native_table():
    """Data Explorer table must render as NATIVE PowerPoint DrawingML table."""
    prs, slide = _make_dummy_slide()
    slide_spec = _make_dummy_slide_spec()

    item = PresentationContentItem(
        content_id="CONTENT-TAB-01",
        role=ContentRole.KEY_FINDING,
        topic_category="Exploration",
        visual_id="TOPIC-TAB-01",
        visual_strategy=VisualStrategy.REUSE_EXISTING,
        visual_type="table",
        source_surface="data_explorer",
        kind=ContentKind.TABLE,
        title="Data Explorer Empirical Table Sample",
        narrative_takeaway="Exploratory table sample",
        table_data=[
            {"State": "Uttar Pradesh", "Population": "199,812,341", "Growth": "20.2%"},
            {"State": "Maharashtra", "Population": "112,374,333", "Growth": "16.0%"},
        ],
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_TAB_01",
        object_name="table_object",
        content_id=item.content_id,
        name="Table Object",
        kind=ContentKind.TABLE,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.NATIVE_REQUIRED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata=item.model_dump(),
    )

    adapter = TableRendererAdapter()
    shape, record = adapter.render(
        slide=slide,
        obj_layout=obj_layout,
        slide_spec=slide_spec,
        content_item=item,
        slide_w_emu=int(prs.slide_width),
        slide_h_emu=int(prs.slide_height),
    )

    assert shape is not None
    assert getattr(shape, "has_table", False) is True
    assert record.render_mode == RenderMode.NATIVE
    assert record.rendered_visual_family == "table"
    assert record.shape_type == ShapeType.TABLE
    assert record.is_native_editable is True
    assert record.source_visual_id == "TOPIC-TAB-01"


def test_hard_rendering_guard_rejects_family_substitution():
    """Hard rendering guard must reject substituting visual family when not PRESERVED_IMAGE."""
    from app.services.presentation.layer_zone_models import LayeredDeckStructureSpec
    from app.services.presentation.layout_models import ResolvedDeckLayout, ResolvedSlideLayout
    from app.services.presentation.render_models import (
        RenderManifest,
        RenderedObjectRecord,
        RenderedSlideRecord,
    )

    validator = RendererValidator()

    slide_spec = _make_dummy_slide_spec()
    deck_structure = LayeredDeckStructureSpec(
        manifest_id="M_TEST",
        manifest_version="1.0",
        slides=[slide_spec],
        object_registry={},
    )

    obj_layout = ResolvedObjectLayout(
        object_id="SLIDE_01__OBJ_SCAT_01",
        object_name="scatter_object",
        content_id="CONTENT-SCAT-01",
        name="Scatter Object",
        kind=ContentKind.CHART,
        binding_role=ContentBindingRole.PRIMARY,
        editability=EditabilityPolicy.NATIVE_REQUIRED,
        bounds=Box2D(left=0.1, top=0.15, width=0.8, height=0.7),
        metadata={
            "visual_strategy": "REUSE_EXISTING",
            "visual_type": "scatter",
        },
    )

    b = Box2D(left=0.05, top=0.05, width=0.9, height=0.9)
    resolved_slide = ResolvedSlideLayout(
        physical_slide_id=slide_spec.physical_slide_id,
        logical_slide_id="LOG_01",
        sequence_index=1,
        zones=[],
        safe_area=ResolvedSafeArea(
            safe_area_bounds=b,
            frame_header_bounds=b,
            frame_footer_bounds=b,
            body_bounds=b,
        ),
        object_layouts=[obj_layout],
    )
    resolved_deck = ResolvedDeckLayout(
        deck_id="DECK_TEST",
        manifest_id="M_TEST",
        manifest_version="1.0",
        slides=[resolved_slide],
        total_objects=1,
    )

    # Record where scatter was illegally substituted with 'bar' while claiming NATIVE
    illegal_record = RenderedObjectRecord(
        object_id=obj_layout.object_id,
        object_name=obj_layout.object_name,
        zone_id="ZONE_01",
        physical_slide_id=slide_spec.physical_slide_id,
        shape_id=1,
        shape_type=ShapeType.CHART,
        is_native_editable=True,
        editability_policy=EditabilityPolicy.NATIVE_REQUIRED,
        applied_bounds=obj_layout.bounds,
        render_mode=RenderMode.NATIVE,
        visual_type="scatter",
        rendered_visual_family="bar",  # ILLEGAL SUBSTITUTION!
        metadata={
            "visual_strategy": "REUSE_EXISTING",
            "visual_type": "scatter",
            "rendered_visual_family": "bar",
        },
    )

    rendered_slide = RenderedSlideRecord(
        physical_slide_id=slide_spec.physical_slide_id,
        logical_slide_id="LOG_01",
        slide_index=1,
        objects=[illegal_record],
        slide_fingerprint="fp123",
    )

    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    cd = CategoryChartData()
    cd.categories = ["A", "B"]
    cd.add_series("S", (1, 2))
    sh = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), cd)
    sh.name = obj_layout.object_name

    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tf:
        pptx_path = tf.name
    prs.save(pptx_path)

    manifest = RenderManifest(
        manifest_id="MAN_TEST",
        deck_id="DECK_TEST",
        pptx_path=pptx_path,
        slides=[rendered_slide],
        total_slides=1,
        total_objects_rendered=1,
        semantic_render_fingerprint="render_fp",
    )

    with pytest.raises(RendererValidationError, match="Visual family substitution prohibited"):
        validator.validate_manifest_and_package(
            manifest=manifest,
            resolved_deck=resolved_deck,
            deck_structure=deck_structure,
        )
