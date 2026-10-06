"""The reference theme must survive selection, pagination and native export."""
import copy

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from app.core import config
from app.services.presentation.slide_layout import resolve_slide
from app.services.presentation.visual.design_tokens import slide_theme_preset
from app.services.report_generator import export_spec_to_pptx, export_spec_to_pdf

pytestmark = pytest.mark.presentation

BRUSH_THEMES = ("amber_brush", "teal_brush", "coral_brush")


def checklist_slide(count=8):
    return {'title': 'Best practices for sick leave', 'layout': 'comparison_split',
            'bullets': [f'Practice {i}: Review recorded absences and confirm the reporting policy.' for i in range(count)],
            'source_label': 'Illustrative practices'}


@pytest.mark.parametrize("theme_id", BRUSH_THEMES)
def test_theme_selection_and_preview_use_the_same_theme_geometry(theme_id):
    client = TestClient(app)
    assert theme_id in {t['id'] for t in client.get('/api/presentations/themes').json()['themes']}
    slide = checklist_slide()
    original = copy.deepcopy(slide)
    preview = client.post(f'/api/presentations/layout-preview?theme_id={theme_id}', json=slide)
    assert preview.status_code == 200
    assert preview.json()['plan'] == resolve_slide(slide, theme_id)
    assert resolve_slide(slide, 'clean_light') != preview.json()['plan']
    assert slide == original
    response = client.get(f'/api/presentations/theme-assets/{theme_id}/background')
    assert response.status_code == 200 and response.content.startswith(b'\x89PNG')
    assert client.get(f'/api/presentations/theme-assets/{theme_id}/font').status_code == 200
    assert client.get('/api/presentations/theme-assets/unknown/background').status_code == 404


@pytest.mark.parametrize("theme_id", BRUSH_THEMES)
def test_checklist_pagination_preserves_all_items_at_readable_sizes(theme_id):
    plan = resolve_slide(checklist_slide(20), theme_id)
    items = [b for page in plan['pages'] for b in page if b['kind'] == 'checklist']
    assert len(plan['pages']) > 1
    assert [item['title'] for item in items] == [f'Practice {i}' for i in range(20)]
    assert all(item['size'] >= 14 and item['y'] >= 190 and item['y']+item['height'] <= 486 for item in items)
    first = [b for b in plan['pages'][0] if b['kind'] == 'checklist']
    assert len({b['x'] for b in first}) == 4
    assert len({b['y'] for b in first}) == 2


@pytest.mark.parametrize("theme_id", BRUSH_THEMES)
def test_export_has_bundled_artwork_editable_checklists_charts_and_tables(theme_id, tmp_path, monkeypatch):
    monkeypatch.setattr(config, 'EXPORTS_DIR', tmp_path)
    slides = [checklist_slide(), {'title': 'Recorded sick leave', 'layout': 'chart_narrative',
              'chart': {'categories': ['Design', 'Operations'], 'series': [{'name': 'Days', 'values': [2, 5]}]}},
              {'title': 'Leave records', 'layout': 'table_detail',
               'table': {'headers': ['Department', 'Days'], 'rows': [['Design', 2], ['Operations', 5]]}}]
    deck = {'id': f'{theme_id}-test', 'metadata': {'theme_id': theme_id}, 'slides': slides}
    prs = Presentation(export_spec_to_pptx(deck))
    theme = slide_theme_preset(theme_id)
    assert len(prs.slides) == 3
    for slide in prs.slides:
        assert slide.shapes[0].name == f"Theme background: {theme['name']}"
        from app.services.presentation.theme_background import theme_background_path
        assert slide.shapes[0].image.blob == theme_background_path(theme_id).read_bytes()
        header = next(shape for shape in slide.shapes if shape.name == 'Text: title')
        run = header.text_frame.paragraphs[0]
        assert str(run.font.color.rgb) == theme['header_text'][1:]
        assert run.font.name == 'Comic Sans MS'
    text = '\n'.join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert all(f'Practice {i}' in text for i in range(8))
    assert len([s for s in prs.slides[0].shapes if s.name == 'Checklist icon']) == 8
    chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
    assert list(chart.series[0].values) == [2, 5]
    table = next(s.table for s in prs.slides[2].shapes if s.has_table)
    assert table.cell(2, 0).text == 'Operations'
    assert export_spec_to_pdf(deck).stat().st_size > 1000


@pytest.mark.parametrize("theme_id", BRUSH_THEMES)
@pytest.mark.parametrize("explicit_theme", [True, False])
def test_copilot_selects_requested_brush_theme(theme_id, explicit_theme):
    from app.routers.copilot import parse_slide_mutation_intent
    prompt = f"Use {theme_id.replace('_', ' ')}" + (" theme" if explicit_theme else "")
    action, params, _ = parse_slide_mutation_intent(prompt)
    assert action == 'change_theme'
    assert params == {'theme_id': theme_id}
