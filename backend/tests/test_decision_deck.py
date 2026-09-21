import copy
import json
from contextlib import closing

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from app.main import app
from app.db.database import get_connection
from app.services.decision_intelligence import build_decision_brief
from app.services.presentation.scope_detector import collect_workspace_evidence
from app.services.presentation.deck_generator import generate_presentation_deck_spec
from app.services.presentation.claim_verifier import verify_presentation_claims
from app.services.presentation.pipeline_orchestrator import execute_presentation_pipeline_async
from app.services.presentation.job_manager import PresentationJobManager
from app.services.report_generator import export_spec_to_pptx


def seed():
    client=TestClient(app)
    content='Region,Date,Revenue,Cost\n'+'\n'.join(f'{"East" if i<15 else "West"},2026-07-{i%28+1:02},{100+i*10},{50+i*2}' for i in range(30))
    assert client.post('/api/upload/file',files={'file':('revenue.csv',content.encode(),'text/csv')}).status_code==200
    return client.get('/api/sheets').json()['sheets'][0]['id']


def make_deck():
    sid=seed(); scope={'scope_type':'single_sheet','sheet_id':sid,'deck_style':'decision_brief','theme_id':'executive_dark'}
    with closing(get_connection()) as conn:
        overview=build_decision_brief(conn,sid)
        evidence=collect_workspace_evidence(conn,scope)
    deck=generate_presentation_deck_spec(scope,evidence['primary_ctx'],evidence)
    return deck,overview,scope


def test_overview_and_presentation_share_snapshot_findings_and_values():
    deck,overview,_=make_deck()
    assert deck['metadata']['snapshot_hash']==overview['snapshot']
    assert deck['decision_brief']['findings']==overview['findings']
    assert verify_presentation_claims(deck,deck['evidence_ledger'])['status']=='PASSED'
    for f in overview['findings']:
        slides=[s for s in deck['slides'] if s['finding_id']==f['id']]
        assert slides and f['observation'] in slides[0]['speaker_notes']
        if f['detail'].get('groups'):
            expected=sorted(g['value'] for g in f['detail']['groups'] if g['value'] is not None)
            actual=sorted(v for s in slides for v in s['chart']['series'][0]['values'])
            assert actual==expected


@pytest.mark.parametrize('field',['title','narrative','bullets','chart','speaker_notes'])
def test_modified_claims_fail_verification_and_export(field):
    deck,_,_=make_deck()
    slide=next(s for s in deck['slides'] if s.get('chart'))
    if field=='chart':slide['chart']['series'][0]['values'][0]*=-1
    elif field=='bullets':slide[field]=['Invented intervention saves 999 dollars']
    else:slide[field]='Invented claim 999'
    assert verify_presentation_claims(deck,deck['evidence_ledger'])['status']=='DISCREPANCIES_FLAGGED'
    with pytest.raises(ValueError,match='changed after evidence'):export_spec_to_pptx(deck)


def test_native_export_retains_editable_chart_values():
    deck,_,_=make_deck();path=export_spec_to_pptx(deck);prs=Presentation(path)
    assert len(prs.slides)==len(deck['slides'])
    charts=[shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    expected=[s['chart'] for s in deck['slides'] if s.get('chart')]
    assert charts and len(charts)==len(expected)
    for chart,spec in zip(charts,expected):
        assert list(chart.series[0].values)==spec['series'][0]['values']


def test_background_pipeline_persists_verified_decision_deck(monkeypatch):
    sid=seed();scope={'scope_type':'single_sheet','sheet_id':sid,'deck_style':'decision_brief'}
    manager=PresentationJobManager();jid=manager.create_job(scope)
    monkeypatch.setattr('time.sleep',lambda _:None)
    execute_presentation_pipeline_async(jid,scope,manager=manager)
    with closing(get_connection()) as conn:
        job=conn.execute('SELECT status,error,deck_id FROM presentation_jobs WHERE id=?',(jid,)).fetchone()
        assert job['status']=='ready',job['error']
        deck=json.loads(conn.execute('SELECT spec_json FROM presentation_decks WHERE id=?',(job['deck_id'],)).fetchone()[0])
    assert deck['metadata']['deck_style']=='decision_brief'
    assert deck['metadata']['validation_summary']['status']=='PASSED'
