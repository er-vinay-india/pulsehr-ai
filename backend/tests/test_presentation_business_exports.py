from contextlib import closing
import json
import zipfile

from fastapi.testclient import TestClient
from pptx import Presentation
import pytest

from app.main import app
from app.core import config
from app.db.database import get_connection
from app.services.presentation.hr_report import build_hr_report
from app.services.presentation.job_manager import PresentationJobManager
from app.services.presentation.pipeline_orchestrator import execute_presentation_pipeline_async
from app.services.presentation.content_validation import semantic_issues
from app.services.report_generator import export_spec_to_pdf, export_spec_to_pptx
from test_presentation_hr_report import fixture_context


@pytest.mark.parametrize('theme', ['clean_light', 'executive_dark', 'amber_brush'])
def test_business_exports_keep_native_values_tables_notes_and_theme(theme):
    deck = build_hr_report({'theme_id': theme}, fixture_context())
    pptx = export_spec_to_pptx(deck)
    prs = Presentation(pptx)
    expected = [s['chart'] for s in deck['slides'] if s.get('chart')]
    charts = [shape.chart for s in prs.slides for shape in s.shapes if shape.has_chart]
    assert len(charts) == len(expected)
    for chart, spec in zip(charts, expected):
        assert list(chart.series[0].values) == spec['series'][0]['values']
        assert chart.chart_title.text_frame.text == spec['title']
        assert str(chart.chart_title.text_frame.paragraphs[0].font.color.rgb) == deck['theme']['primary_text'].lstrip('#')
        assert chart.value_axis.tick_labels.number_format == ('0' if spec['unit'] in {'rows', 'count', 'employees', 'personnel'} else '0.00')
    assert any(shape.has_table for s in prs.slides for shape in s.shapes)
    public = ' '.join(shape.text for s in prs.slides for shape in s.shapes if shape.has_text_frame)
    assert 'HR-001' not in public and 'SQL' not in public
    assert 'Variance Mitigation' not in public and 'Automated governance' not in public
    assert all('snapshot hash' not in s.notes_slide.notes_text_frame.text for s in prs.slides)
    with zipfile.ZipFile(pptx) as archive:
        from lxml import etree
        for name in archive.namelist():
            if name.startswith('ppt/charts/chart') and name.endswith('.xml'):
                root = etree.fromstring(archive.read(name))
                ids = root.xpath('//*[local-name()="axId" or local-name()="crossAx"]/@val')
                assert all(0 <= int(v) < 2**32 for v in ids)
    pdf = export_spec_to_pdf(deck)
    assert pdf.stat().st_size > 1000
    assert not semantic_issues(deck)


@pytest.mark.parametrize('exporter', [export_spec_to_pdf, export_spec_to_pptx])
def test_export_rechecks_changed_notes_and_values(exporter):
    deck = build_hr_report({}, fixture_context())
    deck['slides'][0]['speaker_notes'] = 'The intervention guarantees improved ROI.'
    with pytest.raises(ValueError, match='changed after evidence'):
        exporter(deck)


@pytest.mark.parametrize('theme', ['clean_light', 'amber_brush'])
def test_actual_attendance_job_reaches_ready_and_persists_report(monkeypatch, theme):
    client = TestClient(app)
    csv = 'Emp ID,Department,1st to 5th July 2026,6th to 12th July 2026,Total Attendance,Approved Leaves,Final Attendance\nLB001,Design,0,3,3,2,1\nLB002,Operations,2,4,6,1,5\n'
    response = client.post('/api/upload/file', files={'file': ('wfo.csv', csv.encode(), 'text/csv')})
    assert response.status_code == 200
    sid = client.get('/api/sheets').json()['sheets'][0]['id']
    scope = {'scope_type': 'single_sheet', 'sheet_id': sid, 'theme_id': theme, 'deliverable': 'both',
             'objective': 'Attendance review', 'audience': 'HR managers'}
    manager = PresentationJobManager()
    jid = manager.create_job(scope)
    monkeypatch.setattr('time.sleep', lambda _: None)
    execute_presentation_pipeline_async(jid, scope, manager=manager)
    job = manager.get_job(jid)
    assert job['status'] == 'ready', job.get('error')
    with closing(get_connection()) as conn:
        deck = json.loads(conn.execute('SELECT spec_json FROM presentation_decks WHERE id=?', (job['deck_id'],)).fetchone()[0])
    assert deck['metadata']['content_contract'] == 'hr_attendance_v1'
    assert deck['metadata']['validation_summary']['status'] == 'PASSED'
    assert (config.EXPORTS_DIR/deck['pptx_filename']).exists()
    assert (config.EXPORTS_DIR/deck['pdf_filename']).exists()
    assert deck['slides'][0]['title'] == 'Office attendance review'
