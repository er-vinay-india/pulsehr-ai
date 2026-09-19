import sqlite3
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import app
from app.services import copilot_tools as tools
from app.services import report_generator, sheet_catalog


@pytest.fixture
def source(tmp_path, monkeypatch):
    path = tmp_path / 'test.db'
    def connection():
        conn = sqlite3.connect(path)
        conn.row_factory = sqlite3.Row
        return conn
    conn = connection()
    conn.executescript((Path(__file__).parents[1] / 'app/db/schema.sql').read_text())
    conn.execute("INSERT INTO dataset_uploads(id,filename,original_name,file_type) VALUES (1,'absence.csv','absence.csv','csv')")
    conn.commit()
    conn.close()
    pd.DataFrame({'Absent days': [1] * 350 + [None], 'Department': ['HR'] * 200 + ['IT'] * 151}).to_csv(tmp_path / 'absence.csv', index=False)
    conn = connection()
    sheet_catalog.insert_sheets(conn, 1, sheet_catalog.prepare_sheets(sheet_catalog.read_sheets(tmp_path / 'absence.csv'), 'absence.csv', embed=False))
    conn.commit()
    conn.close()
    monkeypatch.setattr(sheet_catalog, 'get_connection', connection)
    monkeypatch.setattr(tools, 'get_connection', connection)
    monkeypatch.setattr(tools.config, 'UPLOADS_DIR', tmp_path)
    monkeypatch.setattr(report_generator.config, 'EXPORTS_DIR', tmp_path)
    return tmp_path


def test_full_file_sum_beyond_search_index(source):
    result = tools.calculate(tools.CalculationRequest(dataset_id=1, operation='sum', column='Absent days'))
    assert result['source_rows'] == 351
    assert result['results'][0] == {'group': 'All selected rows', 'value': 350., 'rows': 351, 'used_rows': 350, 'missing_rows': 1}


def test_grouped_and_filtered_results(source):
    result = tools.calculate(tools.CalculationRequest(dataset_id=1, operation='sum', column='Absent days', group_by='Department'))
    assert [r['value'] for r in result['results']] == [200., 150.]
    filtered = tools.calculate(tools.CalculationRequest(dataset_id=1, operation='count', filter_column='Department', filter_value='hr'))
    assert filtered['matched_rows'] == 200
    assert filtered['results'][0]['value'] == 200


@pytest.mark.parametrize('operation, expected', [('mean', 1.), ('median', 1.), ('min', 1.), ('max', 1.)])
def test_numeric_operations(source, operation, expected):
    assert tools.calculate(tools.CalculationRequest(dataset_id=1, column='Absent days', operation=operation))['results'][0]['value'] == expected


def test_no_match_and_non_numeric(source):
    result = tools.calculate(tools.CalculationRequest(dataset_id=1, column='Absent days', operation='sum', filter_column='Department', filter_value='Missing'))
    assert result['results'][0]['value'] is None
    with pytest.raises(ValueError, match='non-numeric'):
        tools.calculate(tools.CalculationRequest(dataset_id=1, column='Department', operation='sum'))
    with pytest.raises(ValueError, match='Unknown column'):
        tools.calculate(tools.CalculationRequest(dataset_id=1, column='salary'))


def test_excel_requires_sheet_selection(source):
    with pd.ExcelWriter(source / 'absence.xlsx') as writer:
        pd.DataFrame({'Days': [2, 3]}).to_excel(writer, sheet_name='January', index=False)
        pd.DataFrame({'Days': [4, 5]}).to_excel(writer, sheet_name='February', index=False)
    conn = tools.get_connection()
    conn.execute("UPDATE dataset_uploads SET filename='absence.xlsx'")
    conn.execute('DELETE FROM sheets')
    sheet_catalog.insert_sheets(conn, 1, sheet_catalog.prepare_sheets(sheet_catalog.read_sheets(source / 'absence.xlsx'), 'absence.xlsx', embed=False))
    conn.commit()
    conn.close()
    with pytest.raises(ValueError, match='Choose a source file and sheet'):
        tools.calculate(tools.CalculationRequest(dataset_id=1, column='Days', operation='sum'))
    assert tools.calculate(tools.CalculationRequest(dataset_id=1, sheet='February', column='Days', operation='sum'))['results'][0]['value'] == 9


@pytest.mark.parametrize('expression', ['__import__("os")', '1/0', '2**100', '1e999', '[1]', 'True', '2%3'])
def test_unsafe_or_invalid_arithmetic(expression):
    with pytest.raises(ValueError):
        tools.arithmetic(expression)


def test_arithmetic_and_conservative_routing():
    assert tools.arithmetic('(12 + 8) / 4 * 3') == 15
    assert tools.infer_tool('Average attendance by department').calculation.group_by == 'department'
    assert tools.infer_tool('Average attendance for HR in January') is None
    assert tools.infer_tool('Do not create a presentation') is None


def test_chat_calculation_and_download(source):
    client = TestClient(app)
    request = {'query': 'Calculate absences', 'tool': {'name': 'calculate', 'calculation': {'dataset_id': 1, 'column': 'Absent days', 'operation': 'sum'}}}
    response = client.post('/api/copilot/query', json=request)
    assert response.status_code == 200
    assert response.json()['calculation']['results'][0]['value'] == 350
    request['tool']['name'] = 'presentation'
    response = client.post('/api/copilot/query', json=request)
    artifact = response.json()['artifacts'][0]
    assert client.get(artifact['url']).status_code == 200
    deck = Presentation(source / artifact['name'])
    assert len(deck.slides) == 2
    assert deck.slides[0].shapes[2].table.cell(1, 1).text == '350'
    assert client.get('/api/reports/presentation/files/not-a-report.pptx').status_code == 404
    assert client.post('/api/copilot/query', json={'query': 'x', 'tool': {'name': 'shell'}}).status_code == 422


def test_presentation_uses_live_data_and_unique_names(source, monkeypatch):
    first = report_generator.generate_pptx_presentation()
    second = report_generator.generate_pptx_presentation()
    assert first != second
    deck = Presentation(first)
    text = ' '.join(cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
    assert 'Absent days' in text and 'Department' in text
    assert '92.4%' not in text


@pytest.mark.parametrize('query', [
    'Total overtime in January for HR',
    'How many days was Sofia Sharma absent?',
    'What does average attendance mean?',
    'Explain our total rewards policy',
    'Compare the maximum and minimum attendance for HR',
])
def test_questions_reach_grounded_chat_instead_of_fixed_reply(monkeypatch, query):
    from app.services import ai_copilot
    monkeypatch.setattr(ai_copilot, 'hybrid_search', lambda *args, **kwargs: [])
    monkeypatch.setattr(ai_copilot, 'get_aggregate_context', lambda: 'Verified HR context')
    monkeypatch.setattr(ai_copilot, 'get_available_models', lambda: [])
    calls = []
    class ModelResponse:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            return {'response': 'A response addressing: ' + query}
    class ModelClient:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def post(self, url, json):
            calls.append(json)
            return ModelResponse()
    monkeypatch.setattr(ai_copilot.httpx, 'Client', ModelClient)
    response = TestClient(app).post('/api/copilot/query', json={'query': query})
    assert response.status_code == 200
    assert response.json()['answer'] == 'A response addressing: ' + query
    assert len(calls) == 1
    assert query in calls[0]['prompt']
    assert 'Verified HR context' in calls[0]['prompt']


def test_supported_shortcut_still_executes_without_llm(monkeypatch):
    from app.services import ai_copilot
    def unexpected(*args, **kwargs):
        pytest.fail('A supported tool request should not call the model')
    monkeypatch.setattr(ai_copilot, 'get_available_models', unexpected)
    response = TestClient(app).post('/api/copilot/query', json={'query': 'Calculate (12 + 8) / 4'})
    assert response.json()['tool_used'] == 'arithmetic'
    assert '**5**' in response.json()['answer']


def test_presentation_paginates_groups(source):
    result = tools.calculate(tools.CalculationRequest(dataset_id=1, column='Absent days', operation='sum'))
    result['results'] = [{'group': f'Department {n}', 'value': n, 'used_rows': 1, 'missing_rows': 0} for n in range(19)]
    deck = Presentation(report_generator.generate_calculation_presentation(result))
    assert len(deck.slides) == 4
    assert [len(deck.slides[n].shapes[2].table.rows) for n in range(3)] == [9, 9, 4]
